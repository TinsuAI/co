from __future__ import annotations

import importlib.util
import json
import sys
import tempfile
import unittest
from dataclasses import dataclass
from datetime import date, timedelta
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
SCRIPTS_DIR = ROOT / "scripts"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

SPEC = importlib.util.spec_from_file_location(
    "growatt_replacement_runner_test",
    SCRIPTS_DIR / "growatt_replacement_runner.py",
)
if SPEC is None or SPEC.loader is None:
    raise RuntimeError("Unable to load growatt_replacement_runner.py for tests")
runner = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = runner
SPEC.loader.exec_module(runner)


class FakeBaseline:
    TARGET_RVC = 35.0

    @staticmethod
    def is_vietnam_origin(origin: str | None) -> bool:
        return str(origin or "").strip().upper() in {"VIETNAM", "VN"}

    @staticmethod
    def is_bucket_eligible_for_export(bucket, export_line, import_lead_days: int, max_import_age_days: int) -> bool:
        if bucket.import_date is None:
            return False
        if bucket.import_date > export_line.export_date - timedelta(days=import_lead_days):
            return False
        if max_import_age_days > 0 and bucket.import_date < export_line.export_date - timedelta(days=max_import_age_days):
            return False
        return True


@dataclass
class FakeBomLine:
    material_code: str
    qty_per_unit: float


@dataclass
class FakeVariant:
    variant_id: str
    bom_code: str
    lines: list[FakeBomLine]


@dataclass
class FakeExportLine:
    model_code: str
    declaration_no: str
    declaration_item_no: int
    export_date: date
    quantity: float
    unit_price_usd: float
    invoice_no: str = ""


def make_bucket(
    *,
    basis: str,
    material: str,
    bucket_id: str,
    qty: float,
    unit_price: float,
    import_date: date,
    candidate_class: str = "confirmed_clean",
    mapping_status: str | None = None,
    mapping_confidence: str = "medium",
    variant_scope: tuple[str, ...] = (),
    origin: str = "CHINA",
) -> runner.ReplacementBasisBucket:
    if mapping_status is None:
        if candidate_class == "candidate_clean":
            mapping_status = "candidate_exact_dm_match"
        elif candidate_class == "ambiguity_review":
            mapping_status = "candidate_not_in_dm"
        else:
            mapping_status = "confirmed_declared_equals_dm_code"
    return runner.ReplacementBasisBucket(
        custom_code_basis=basis,
        erp_material_code=material,
        bucket_id=bucket_id,
        tracking_key=bucket_id,
        declaration_no="DECL",
        declaration_item_no=1,
        import_date=import_date,
        source="BCCT",
        source_row_no="1",
        hs_code="HS",
        name=material,
        origin=origin,
        unit_price_usd=unit_price,
        exchange_rate=1.0,
        remaining_qty=qty,
        candidate_class=candidate_class,
        mapping_status=mapping_status,
        mapping_confidence=mapping_confidence,
        evidence_source="mapping",
        shipment_variant_scope="confirmed",
        variant_scope=variant_scope,
        variant_hit_count=len(variant_scope),
        candidate_material_code=material,
    )


class GrowattReplacementRunnerTests(unittest.TestCase):
    def setUp(self) -> None:
        self.baseline = FakeBaseline()
        self.export_line = FakeExportLine(
            model_code="MODEL-A",
            declaration_no="DECL-1",
            declaration_item_no=1,
            export_date=date(2026, 4, 10),
            quantity=10.0,
            unit_price_usd=10.0,
        )
        self.variant = FakeVariant(
            variant_id="MODEL-A__block1",
            bom_code="MODEL-A",
            lines=[FakeBomLine(material_code="ORIG", qty_per_unit=1.0)],
        )

    def test_choose_custom_code_basis_prefers_shared_declared_group(self) -> None:
        buckets = [
            make_bucket(
                basis="ORIG",
                material="ORIG",
                bucket_id="orig-exact",
                qty=10.0,
                unit_price=5.0,
                import_date=date(2026, 4, 1),
            ),
            make_bucket(
                basis="TUDIEN",
                material="ORIG",
                bucket_id="orig-family",
                qty=5.0,
                unit_price=5.0,
                import_date=date(2026, 4, 1),
            ),
            make_bucket(
                basis="TUDIEN",
                material="REPL",
                bucket_id="repl-family",
                qty=100.0,
                unit_price=1.0,
                import_date=date(2026, 4, 1),
            ),
        ]
        _, stock_by_basis, _ = runner.clone_stock_indexes(buckets)
        stats = runner.build_material_basis_stats(buckets)
        chosen = runner.choose_custom_code_basis("ORIG", stats, stock_by_basis)
        self.assertEqual(chosen, "TUDIEN")

    def test_build_replacement_basis_uses_label_code_for_nk2_internal_declared_code(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            tmp_path = Path(tmpdir)
            shared_normalized_dir = tmp_path / "shared"
            shipment_dir = tmp_path / "b282"
            shared_normalized_dir.mkdir(parents=True, exist_ok=True)
            (shipment_dir / "normalized").mkdir(parents=True, exist_ok=True)

            (shipment_dir / "normalized" / "b282-variant-admissibility.csv").write_text(
                "tracking_key,admissible_material_code,variant_id,shipment_variant_scope\n",
                encoding="utf-8",
            )
            (shared_normalized_dir / "co-stock-tracking-updated.csv").write_text(
                "\n".join(
                    [
                        "source,tracking_key,declaration_no,declaration_date,declaration_item_no,declared_code,lookup_material_code,confirmed_lookup_code,final_lookup_key,matched_nk2_source_row_no,code_extraction_status,dm_match_status,mapping_status,lookup_confidence,hs_code,label_code,paren_code_candidates,name,description_clean,origin,source_unit_price,unit_price_usd,tax_unit_price,price_normalization_basis,import_qty,used_qty,remaining_qty,unit,partner_name,invoice_no,invoice_date,exchange_rate,source_exchange_rate,nk2_source_row_no,bcct_source_row_no",
                        "NK2_ONLY,track-1,DECL,2026-04-01,1,007.0030600,007.0030600,007.0030600,,1,fallback_declared_code,confirmed_exact_dm_match,confirmed_declared_equals_dm_code,medium,85423900,IC,,IC#&Some IC part (007.0030600),Some IC part,CHINA,1,1,1,source_price,10,0,10,PIECES,SUPPLIER,INV,2026-04-01,1,1,1,",
                        "NK2_ONLY,track-2,DECL,2026-04-01,1,001.0036000,001.0036000,001.0036000,,2,fallback_declared_code,confirmed_exact_dm_match,confirmed_declared_equals_dm_code,medium,85332100,DIENTRO.CHIP,,DIENTRO.CHIP#&Chip resistor (001.0036000),Chip resistor,CHINA,1,1,1,source_price,10,0,10,PIECES,SUPPLIER,INV,2026-04-01,1,1,2,",
                    ]
                ),
                encoding="utf-8",
            )

            buckets = runner.build_replacement_basis(shared_normalized_dir, shipment_dir)

        self.assertEqual(len(buckets), 2)
        by_material = {bucket.erp_material_code: bucket.custom_code_basis for bucket in buckets}
        self.assertEqual(by_material["007.0030600"], "IC")
        self.assertEqual(by_material["001.0036000"], "DIENTRO.CHIP")

    def test_build_candidate_rows_include_clean_date_blocked_and_ambiguity_with_flags(self) -> None:
        material = runner.MaterialState(
            original_material_code="ORIG",
            need_qty=10.0,
            exact_allocated_qty=0.0,
            unmet_qty=10.0,
            custom_code_basis="TUDIEN",
        )
        stock_by_basis = {
            "TUDIEN": [
                make_bucket(
                    basis="TUDIEN",
                    material="REPL-CLEAN",
                    bucket_id="clean",
                    qty=10.0,
                    unit_price=1.0,
                    import_date=date(2026, 4, 1),
                    variant_scope=(self.variant.variant_id,),
                ),
                make_bucket(
                    basis="TUDIEN",
                    material="REPL-AMBIG",
                    bucket_id="ambig",
                    qty=10.0,
                    unit_price=1.0,
                    import_date=date(2026, 4, 1),
                    candidate_class="ambiguity_review",
                ),
                make_bucket(
                    basis="TUDIEN",
                    material="REPL-LATE",
                    bucket_id="late",
                    qty=10.0,
                    unit_price=1.0,
                    import_date=date(2026, 4, 9),
                ),
            ]
        }
        rows = runner.build_candidate_rows(
            self.baseline,
            runner.CandidateContext(
                starting_point_id="seed",
                sequence_no=1,
                iteration_no=1,
                model_code=self.export_line.model_code,
                variant_id=self.variant.variant_id,
                declaration_no=self.export_line.declaration_no,
            ),
            self.export_line,
            self.variant,
            [material],
            runner.summarize_product(self.baseline, self.export_line, [material]),
            material,
            stock_by_basis,
            {("ORIG", "REPL-CLEAN")},
            import_lead_days=2,
            max_import_age_days=0,
        )
        self.assertEqual(
            {row["candidate_material_code"] for row in rows},
            {"REPL-CLEAN", "REPL-AMBIG", "REPL-LATE"},
        )
        by_code = {row["candidate_material_code"]: row for row in rows}
        self.assertTrue(by_code["REPL-CLEAN"]["staff_known_substitute"])
        self.assertEqual(by_code["REPL-CLEAN"]["review_scope"], "clean_admissible")
        self.assertEqual(by_code["REPL-CLEAN"]["risk_flags"], "")
        self.assertEqual(by_code["REPL-AMBIG"]["review_scope"], "ambiguity_review")
        self.assertEqual(by_code["REPL-AMBIG"]["risk_flags"], "ambiguity")
        self.assertTrue(by_code["REPL-AMBIG"]["commit_scope_eligible"])
        self.assertEqual(by_code["REPL-LATE"]["review_scope"], "date_blocked")
        self.assertEqual(by_code["REPL-LATE"]["risk_flags"], "date_blocked")
        self.assertFalse(by_code["REPL-LATE"]["commit_scope_eligible"])

    def test_build_material_plan_splits_and_prefers_lower_cost_supply(self) -> None:
        material = runner.MaterialState(
            original_material_code="ORIG",
            need_qty=10.0,
            exact_allocated_qty=0.0,
            unmet_qty=10.0,
            custom_code_basis="TUDIEN",
        )
        stock_by_basis = {
            "TUDIEN": [
                make_bucket(
                    basis="TUDIEN",
                    material="REPL-1",
                    bucket_id="r1",
                    qty=6.0,
                    unit_price=1.0,
                    import_date=date(2026, 4, 1),
                ),
                make_bucket(
                    basis="TUDIEN",
                    material="REPL-2",
                    bucket_id="r2",
                    qty=4.0,
                    unit_price=2.0,
                    import_date=date(2026, 4, 1),
                ),
                make_bucket(
                    basis="TUDIEN",
                    material="REPL-3",
                    bucket_id="r3",
                    qty=10.0,
                    unit_price=5.0,
                    import_date=date(2026, 4, 1),
                ),
            ]
        }
        plan = runner.build_material_plan(
            self.baseline,
            self.export_line,
            self.variant,
            [material],
            runner.summarize_product(self.baseline, self.export_line, [material]),
            material,
            stock_by_basis,
            set(),
            import_lead_days=2,
            max_import_age_days=0,
        )
        self.assertIsNotNone(plan)
        self.assertEqual(plan.covered_qty, 10.0)
        self.assertEqual(
            [item.material_code for item in plan.allocations],
            ["REPL-1", "REPL-2"],
        )

    def test_run_starting_point_can_swap_sufficient_line_to_lower_non_origin_cost_and_pass(self) -> None:
        initial_stock = {
            "ORIG": [
                make_bucket(
                    basis="FAMILY",
                    material="ORIG",
                    bucket_id="orig-expensive",
                    qty=1.0,
                    unit_price=80.0,
                    import_date=date(2026, 4, 1),
                )
            ],
            "ALT-VN": [
                make_bucket(
                    basis="FAMILY",
                    material="ALT-VN",
                    bucket_id="alt-vn",
                    qty=1.0,
                    unit_price=5.0,
                    import_date=date(2026, 4, 1),
                    origin="VIETNAM",
                )
            ],
        }
        basis_stats = {
            "ORIG": {
                "FAMILY": runner.BasisStats(
                    custom_code_basis="FAMILY",
                    distinct_material_codes={"ORIG", "ALT-VN"},
                    total_qty=2.0,
                    confirmed_qty=2.0,
                )
            }
        }
        variant = FakeVariant("MODEL-A__block1", "MODEL-A", [FakeBomLine("ORIG", 1.0)])
        export_line = FakeExportLine("MODEL-A", "DECL-1", 1, date(2026, 4, 10), 1.0, 100.0)

        with tempfile.TemporaryDirectory() as tmpdir:
            result = runner.run_starting_point(
                self.baseline,
                shipment_id="SHIP",
                starting_point_id="seed",
                export_lines=[export_line],
                variant_by_id={"MODEL-A": variant},
                initial_stock_by_material=initial_stock,
                material_basis_stats=basis_stats,
                staff_substitutes=set(),
                import_lead_days=2,
                max_import_age_days=0,
                output_dir=Path(tmpdir),
            )

        product = result["seed_summary"]["product_status"][0]
        self.assertTrue(product["passes_after"])
        self.assertEqual(product["unmet_after"], 0.0)
        self.assertEqual(product["changed_materials"], "ORIG")

    def test_run_starting_point_can_commit_no_lookup_family_supply(self) -> None:
        initial_stock = {
            "UNMAPPED::FAMILY": [
                make_bucket(
                    basis="FAMILY",
                    material="UNMAPPED::FAMILY",
                    bucket_id="family-unmapped",
                    qty=10.0,
                    unit_price=1.0,
                    import_date=date(2026, 4, 1),
                    candidate_class="ambiguity_review",
                    mapping_status="no_lookup_code",
                    origin="VIETNAM",
                )
            ]
        }
        basis_stats = {
            "ORIG": {
                "FAMILY": runner.BasisStats(
                    custom_code_basis="FAMILY",
                    distinct_material_codes={"ORIG", "UNMAPPED::FAMILY"},
                    total_qty=10.0,
                )
            }
        }
        variant = FakeVariant("MODEL-A__block1", "MODEL-A", [FakeBomLine("ORIG", 1.0)])
        export_line = FakeExportLine("MODEL-A", "DECL-1", 1, date(2026, 4, 10), 10.0, 10.0)

        with tempfile.TemporaryDirectory() as tmpdir:
            result = runner.run_starting_point(
                self.baseline,
                shipment_id="SHIP",
                starting_point_id="seed",
                export_lines=[export_line],
                variant_by_id={"MODEL-A": variant},
                initial_stock_by_material=initial_stock,
                material_basis_stats=basis_stats,
                staff_substitutes=set(),
                import_lead_days=2,
                max_import_age_days=0,
                output_dir=Path(tmpdir),
            )

        product = result["seed_summary"]["product_status"][0]
        self.assertTrue(product["passes_after"])
        candidate_rows = result["candidate_rows"]
        self.assertTrue(any(row["candidate_material_code"] == "UNMAPPED::FAMILY" for row in candidate_rows))

    def test_run_starting_point_stops_after_first_pass_to_preserve_later_substitutes(self) -> None:
        initial_stock = {
            "A": [
                make_bucket(
                    basis="FAMILY-A",
                    material="A",
                    bucket_id="a-1",
                    qty=1.0,
                    unit_price=40.0,
                    import_date=date(2026, 4, 1),
                ),
                make_bucket(
                    basis="FAMILY-A",
                    material="A",
                    bucket_id="a-2",
                    qty=1.0,
                    unit_price=40.0,
                    import_date=date(2026, 4, 1),
                ),
            ],
            "B": [
                make_bucket(
                    basis="FAMILY-B",
                    material="B",
                    bucket_id="b-1",
                    qty=1.0,
                    unit_price=40.0,
                    import_date=date(2026, 4, 1),
                ),
                make_bucket(
                    basis="FAMILY-B",
                    material="B",
                    bucket_id="b-2",
                    qty=1.0,
                    unit_price=40.0,
                    import_date=date(2026, 4, 1),
                ),
            ],
            "ALT-A": [
                make_bucket(
                    basis="FAMILY-A",
                    material="ALT-A",
                    bucket_id="alt-a",
                    qty=1.0,
                    unit_price=5.0,
                    import_date=date(2026, 4, 1),
                    origin="VIETNAM",
                )
            ],
            "ALT-B": [
                make_bucket(
                    basis="FAMILY-B",
                    material="ALT-B",
                    bucket_id="alt-b",
                    qty=1.0,
                    unit_price=5.0,
                    import_date=date(2026, 4, 1),
                    origin="VIETNAM",
                )
            ],
        }
        basis_stats = {
            "A": {
                "FAMILY-A": runner.BasisStats(
                    custom_code_basis="FAMILY-A",
                    distinct_material_codes={"A", "ALT-A"},
                    total_qty=3.0,
                    confirmed_qty=3.0,
                )
            },
            "B": {
                "FAMILY-B": runner.BasisStats(
                    custom_code_basis="FAMILY-B",
                    distinct_material_codes={"B", "ALT-B"},
                    total_qty=3.0,
                    confirmed_qty=3.0,
                )
            },
        }
        variant = FakeVariant(
            "MODEL-A__block1",
            "MODEL-A",
            [FakeBomLine("A", 1.0), FakeBomLine("B", 1.0)],
        )
        export_lines = [
            FakeExportLine("MODEL-A", "DECL-1", 1, date(2026, 4, 10), 1.0, 100.0),
            FakeExportLine("MODEL-A", "DECL-2", 2, date(2026, 4, 10), 1.0, 100.0),
        ]

        with tempfile.TemporaryDirectory() as tmpdir:
            result = runner.run_starting_point(
                self.baseline,
                shipment_id="SHIP",
                starting_point_id="seed",
                export_lines=export_lines,
                variant_by_id={"MODEL-A": variant},
                initial_stock_by_material=initial_stock,
                material_basis_stats=basis_stats,
                staff_substitutes=set(),
                import_lead_days=2,
                max_import_age_days=0,
                output_dir=Path(tmpdir),
            )

        products = result["seed_summary"]["product_status"]
        self.assertEqual(len(products), 2)
        self.assertTrue(products[0]["passes_after"])
        self.assertTrue(products[1]["passes_after"])
        self.assertEqual(products[0]["changed_materials"], "A")
        self.assertEqual(products[1]["changed_materials"], "B")

    def test_run_starting_point_keeps_stock_isolated_between_seeds(self) -> None:
        initial_stock = {
            "REPL": [
                make_bucket(
                    basis="BASIS",
                    material="REPL",
                    bucket_id="shared-repl",
                    qty=10.0,
                    unit_price=1.0,
                    import_date=date(2026, 4, 1),
                )
            ]
        }
        basis_stats = {
            "ORIG": {
                "BASIS": runner.BasisStats(
                    custom_code_basis="BASIS",
                    distinct_material_codes={"ORIG"},
                    total_qty=10.0,
                )
            }
        }
        variant_a = FakeVariant("MODEL-A__block1", "MODEL-A", [FakeBomLine("ORIG", 1.0)])
        variant_b = FakeVariant("MODEL-B__block1", "MODEL-B", [FakeBomLine("ORIG", 1.0)])
        export_a = FakeExportLine("MODEL-A", "DECL-A", 1, date(2026, 4, 10), 10.0, 10.0)
        export_b = FakeExportLine("MODEL-B", "DECL-B", 2, date(2026, 4, 10), 10.0, 10.0)

        with tempfile.TemporaryDirectory() as tmpdir:
            output_dir = Path(tmpdir)
            result_a = runner.run_starting_point(
                self.baseline,
                shipment_id="SHIP",
                starting_point_id="seed-a",
                export_lines=[export_a, export_b],
                variant_by_id={"MODEL-A": variant_a, "MODEL-B": variant_b},
                initial_stock_by_material=initial_stock,
                material_basis_stats=basis_stats,
                staff_substitutes=set(),
                import_lead_days=2,
                max_import_age_days=0,
                output_dir=output_dir,
            )
            result_b = runner.run_starting_point(
                self.baseline,
                shipment_id="SHIP",
                starting_point_id="seed-b",
                export_lines=[export_b, export_a],
                variant_by_id={"MODEL-A": variant_a, "MODEL-B": variant_b},
                initial_stock_by_material=initial_stock,
                material_basis_stats=basis_stats,
                staff_substitutes=set(),
                import_lead_days=2,
                max_import_age_days=0,
                output_dir=output_dir,
            )

        summary_a = result_a["seed_summary"]["product_status"]
        summary_b = result_b["seed_summary"]["product_status"]
        self.assertTrue(summary_a[0]["passes_after"])
        self.assertFalse(summary_a[1]["passes_after"])
        self.assertEqual(summary_a[0]["model_code"], "MODEL-A")
        self.assertTrue(summary_b[0]["passes_after"])
        self.assertFalse(summary_b[1]["passes_after"])
        self.assertEqual(summary_b[0]["model_code"], "MODEL-B")

    def test_run_starting_point_rejects_partial_line_coverage(self) -> None:
        initial_stock = {
            "REPL": [
                make_bucket(
                    basis="BASIS",
                    material="REPL",
                    bucket_id="partial-repl",
                    qty=4.0,
                    unit_price=1.0,
                    import_date=date(2026, 4, 1),
                )
            ]
        }
        basis_stats = {
            "ORIG": {
                "BASIS": runner.BasisStats(
                    custom_code_basis="BASIS",
                    distinct_material_codes={"ORIG"},
                    total_qty=4.0,
                )
            }
        }

        with tempfile.TemporaryDirectory() as tmpdir:
            result = runner.run_starting_point(
                self.baseline,
                shipment_id="SHIP",
                starting_point_id="seed",
                export_lines=[self.export_line],
                variant_by_id={self.export_line.model_code: self.variant},
                initial_stock_by_material=initial_stock,
                material_basis_stats=basis_stats,
                staff_substitutes=set(),
                import_lead_days=2,
                max_import_age_days=0,
                output_dir=Path(tmpdir),
            )

        product = result["seed_summary"]["product_status"][0]
        self.assertFalse(product["passes_after"])
        self.assertEqual(product["unmet_after"], product["unmet_before"])
        self.assertEqual(product["changed_materials"], "")

    def test_extend_starting_points_with_heuristics_adds_two_distinct_seeds(self) -> None:
        scenario_existing = {
            "scenario_id": "SCEN-1",
            "stock_sufficient": False,
            "passes_rvc": False,
            "min_margin": -10.0,
            "total_unmet_qty": 10.0,
            "product_results": [
                {
                    "model_code": "MODEL-A",
                    "variant_id": "MODEL-A__block1",
                    "declaration_no": "DECL-1",
                    "materials": [
                        {
                            "material_code": "ORIG",
                            "need_qty": 10.0,
                            "unmet_qty": 0.0,
                            "non_origin_value_usd": 20.0,
                        }
                    ],
                }
            ],
        }
        scenario_unmet = {
            "scenario_id": "SCEN-2",
            "stock_sufficient": False,
            "passes_rvc": False,
            "min_margin": -8.0,
            "total_unmet_qty": 5.0,
            "product_results": [
                {
                    "model_code": "MODEL-A",
                    "variant_id": "MODEL-A__block1",
                    "declaration_no": "DECL-1",
                    "materials": [
                        {
                            "material_code": "ORIG",
                            "need_qty": 10.0,
                            "unmet_qty": 5.0,
                            "non_origin_value_usd": 20.0,
                        }
                    ],
                }
            ],
        }
        scenario_cost = {
            "scenario_id": "SCEN-3",
            "stock_sufficient": True,
            "passes_rvc": False,
            "min_margin": -2.0,
            "total_unmet_qty": 0.0,
            "product_results": [
                {
                    "model_code": "MODEL-A",
                    "variant_id": "MODEL-A__block1",
                    "declaration_no": "DECL-1",
                    "materials": [
                        {
                            "material_code": "ORIG",
                            "need_qty": 10.0,
                            "unmet_qty": 0.0,
                            "non_origin_value_usd": 90.0,
                        }
                    ],
                }
            ],
        }
        starting_points_payload = {
            "heuristic_best": {
                "starting_point_id": "heuristic_best",
                "variant_strategy": "existing",
                "sequence_strategy": "declaration_order",
                "export_order": ["MODEL-A:DECL-1:1"],
                "scenario": scenario_existing,
            }
        }
        stock_by_basis = {
            "FAMILY": [
                make_bucket(
                    basis="FAMILY",
                    material="ALT-VN",
                    bucket_id="alt-vn",
                    qty=20.0,
                    unit_price=1.0,
                    import_date=date(2026, 4, 1),
                    origin="VIETNAM",
                )
            ]
        }
        basis_stats = {
            "ORIG": {
                "FAMILY": runner.BasisStats(
                    custom_code_basis="FAMILY",
                    distinct_material_codes={"ORIG", "ALT-VN"},
                    total_qty=20.0,
                    confirmed_qty=20.0,
                )
            }
        }
        export_line = FakeExportLine("MODEL-A", "DECL-1", 1, date(2026, 4, 10), 10.0, 10.0)

        with tempfile.TemporaryDirectory() as tmpdir:
            scenarios_path = Path(tmpdir) / "baseline-scenarios.json"
            scenarios_path.write_text(
                json.dumps({"SHIP": [scenario_existing, scenario_unmet, scenario_cost]}),
                encoding="utf-8",
            )
            extended, analysis = runner.extend_starting_points_with_heuristics(
                self.baseline,
                "SHIP",
                starting_points_payload,
                scenarios_path,
                [export_line],
                {("MODEL-A", "DECL-1"): export_line},
                stock_by_basis,
                basis_stats,
                2,
                0,
            )

        self.assertIn("replaceability_unmet_best", extended)
        self.assertIn("replaceability_cost_down_best", extended)
        self.assertEqual(extended["replaceability_unmet_best"]["scenario"]["scenario_id"], "SCEN-2")
        self.assertEqual(extended["replaceability_cost_down_best"]["scenario"]["scenario_id"], "SCEN-3")
        selected = {row["selected_as"]: row["scenario_id"] for row in analysis if row.get("selected_as")}
        self.assertEqual(selected["replaceability_unmet_best"], "SCEN-2")
        self.assertEqual(selected["replaceability_cost_down_best"], "SCEN-3")

    def test_write_review_workbook_creates_excel_sheets(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            workbook_path = Path(tmpdir) / "review.xlsx"
            runner.write_review_workbook(
                workbook_path,
                basis_rows=[
                    {
                        "custom_code_basis": "BASIS",
                        "erp_material_code": "REPL",
                        "bucket_id": "b1",
                        "tracking_key": "b1",
                        "candidate_class": "confirmed_clean",
                        "mapping_status": "confirmed_declared_equals_dm_code",
                        "mapping_confidence": "medium",
                        "evidence_source": "mapping",
                        "remaining_qty": 10.0,
                        "unit_price_usd": 1.0,
                        "origin": "CHINA",
                        "shipment_variant_scope": "confirmed",
                        "variant_scope": "",
                        "variant_hit_count": 1,
                        "source": "BCCT",
                        "source_row_no": "1",
                        "declaration_no": "DECL",
                        "declaration_item_no": 1,
                        "import_date": date(2026, 4, 1),
                        "hs_code": "HS",
                        "name": "REPL",
                    }
                ],
                seed_summaries=[
                    {
                        "starting_point_id": "seed",
                        "pass_count": 0,
                        "final_unmet_qty": 1.0,
                        "min_margin": -1.0,
                        "changed_products": ["MODEL-A"],
                        "changed_materials": ["ORIG"],
                    }
                ],
                product_status_rows=[
                    {
                        "starting_point_id": "seed",
                        "sequence_no": 1,
                        "model_code": "MODEL-A",
                        "variant_id": "MODEL-A__block1",
                        "bom_code": "MODEL-A",
                        "declaration_no": "DECL",
                        "rvc_before": 10.0,
                        "rvc_after": 11.0,
                        "stock_before": False,
                        "stock_after": False,
                        "passes_before": False,
                        "passes_after": False,
                        "unmet_before": 2.0,
                        "unmet_after": 1.0,
                        "changed_materials": "ORIG",
                        "snapshot_path": "/tmp/mock.json",
                    }
                ],
                candidate_rows_by_seed={
                    "seed": [
                        {
                            "starting_point_id": "seed",
                            "sequence_no": 1,
                            "iteration_no": 1,
                            "model_code": "MODEL-A",
                            "variant_id": "MODEL-A__block1",
                            "declaration_no": "DECL",
                            "original_material_code": "ORIG",
                            "custom_code_basis": "BASIS",
                            "candidate_rank": 1,
                            "candidate_material_code": "REPL",
                            "admissible_qty": 10.0,
                            "allocated_preview_qty": 1.0,
                            "clean_qty_available": 10.0,
                            "date_blocked_qty_available": 0.0,
                            "ambiguity_qty_available": 0.0,
                            "ambiguity_date_blocked_qty_available": 0.0,
                            "review_scope": "clean_admissible",
                            "risk_flags": "",
                            "risk_flag_count": 0,
                            "commit_scope_eligible": True,
                            "estimate_scope": "strict_clean",
                            "mapping_confidence": "medium",
                            "confidence_rank": 3,
                            "mapping_evidence_source": "mapping",
                            "staff_known_substitute": False,
                            "estimated_rvc_after": 11.0,
                            "estimated_rvc_gain": 1.0,
                            "would_reach_target": False,
                            "realized_cost_usd": 1.0,
                            "variant_hit_scope": "MODEL-A__block1",
                            "candidate_class": "confirmed_clean",
                            "remaining_unmet_after_candidate": 1.0,
                            "bucket_cost_basis": "[]",
                        }
                    ]
                },
            )
            workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
            self.assertEqual(
                workbook.sheetnames,
                ["Seed Summary", "Product Status", "Custom Code Basis", "Cand seed"],
            )

    def test_write_rvc_and_bom_workbooks_create_expected_sheets(self) -> None:
        snapshot = {
            "shipment_id": "SHIP",
            "starting_point_id": "seed",
            "sequence_no": 1,
            "model_code": "MODEL-A",
            "variant_id": "MODEL-A__block1",
            "bom_code": "MODEL-A",
            "declaration_no": "DECL",
            "export_qty": 10.0,
            "export_date": "2026-04-10",
            "rvc_before": 20.0,
            "rvc_after": 40.0,
            "bom_before": {
                "rvc_percent": 20.0,
                "margin_to_threshold": -15.0,
                "stock_sufficient": False,
                "passes_rvc": False,
                "unmet_qty_total": 2.0,
                "non_origin_value_usd": 80.0,
                "materials": [
                    {
                        "original_material_code": "ORIG",
                        "custom_code_basis": "FAMILY",
                        "need_qty": 10.0,
                        "exact_allocated_qty": 8.0,
                        "replacement_allocated_qty": 0.0,
                        "unmet_qty": 2.0,
                        "exact_allocations": [],
                        "replacement_allocations": [],
                    }
                ],
            },
            "accepted_replacement_decisions": [],
            "bom_after": {
                "rvc_percent": 40.0,
                "margin_to_threshold": 5.0,
                "stock_sufficient": True,
                "passes_rvc": True,
                "unmet_qty_total": 0.0,
                "non_origin_value_usd": 60.0,
                "materials": [
                    {
                        "original_material_code": "ORIG",
                        "custom_code_basis": "FAMILY",
                        "need_qty": 10.0,
                        "exact_allocated_qty": 0.0,
                        "replacement_allocated_qty": 10.0,
                        "unmet_qty": 0.0,
                        "exact_allocations": [],
                        "replacement_allocations": [
                            {
                                "bucket_id": "b1",
                                "tracking_key": "t1",
                                "material_code": "REPL",
                                "custom_code_basis": "FAMILY",
                                "declaration_no": "D1",
                                "declaration_item_no": 1,
                                "import_date": "2026-04-01",
                                "allocated_qty": 10.0,
                                "unit_price_usd": 6.0,
                                "exchange_rate": 1.0,
                                "origin": "CHINA",
                                "source": "BCCT",
                                "source_row_no": "1",
                                "candidate_class": "ambiguity_review",
                                "mapping_status": "no_lookup_code",
                                "mapping_confidence": "none",
                                "evidence_source": "no_lookup_code",
                                "staff_known_substitute": False,
                            }
                        ],
                    }
                ],
            },
            "material_mappings": [],
            "local_stock_state_after_product": {},
        }
        export_line = FakeExportLine(
            model_code="MODEL-A",
            declaration_no="DECL",
            declaration_item_no=1,
            export_date=date(2026, 4, 10),
            quantity=10.0,
            unit_price_usd=10.0,
        )
        setattr(export_line, "name", "Model A")
        setattr(export_line, "hs_code", "85044090")

        with tempfile.TemporaryDirectory() as tmpdir:
            tmp = Path(tmpdir)
            rvc_path = tmp / "rvc.xlsx"
            bom_path = tmp / "bom.xlsx"
            runner.write_rvc_standard_workbook(
                rvc_path,
                [snapshot],
                {("MODEL-A", "DECL"): export_line},
                {"REPL": {"name": "Replacement Material", "hs_code": "850440", "unit": "PIECES"}},
            )
            runner.write_bom_before_after_workbook(bom_path, [snapshot])

            rvc_wb = openpyxl.load_workbook(rvc_path, read_only=True, data_only=True)
            bom_wb = openpyxl.load_workbook(bom_path, read_only=True, data_only=True)

        self.assertEqual(rvc_wb.sheetnames, ["01-MODEL-A"])
        self.assertEqual(rvc_wb["01-MODEL-A"]["B3"].value, "BẢNG KÊ KHAI HÀNG HÓA XUẤT KHẨU ĐẠT TIÊU CHÍ “RVC”")
        self.assertEqual(bom_wb.sheetnames, ["01-MODEL-A"])
        self.assertEqual(bom_wb["01-MODEL-A"]["A1"].value, "Seed: seed")

    def test_rvc_and_bom_workbooks_mark_unmet_and_changed_rows(self) -> None:
        snapshot = {
            "shipment_id": "SHIP",
            "starting_point_id": "seed",
            "sequence_no": 1,
            "model_code": "MODEL-A",
            "variant_id": "MODEL-A__block1",
            "bom_code": "MODEL-A",
            "declaration_no": "DECL",
            "export_qty": 10.0,
            "export_date": "2026-04-10",
            "rvc_before": 20.0,
            "rvc_after": 30.0,
            "bom_before": {
                "rvc_percent": 20.0,
                "margin_to_threshold": -15.0,
                "stock_sufficient": False,
                "passes_rvc": False,
                "unmet_qty_total": 2.0,
                "non_origin_value_usd": 80.0,
                "materials": [
                    {
                        "original_material_code": "ORIG",
                        "custom_code_basis": "FAMILY",
                        "need_qty": 10.0,
                        "exact_allocated_qty": 8.0,
                        "replacement_allocated_qty": 0.0,
                        "unmet_qty": 2.0,
                        "exact_allocations": [],
                        "replacement_allocations": [],
                    }
                ],
            },
            "accepted_replacement_decisions": [],
            "bom_after": {
                "rvc_percent": 30.0,
                "margin_to_threshold": -5.0,
                "stock_sufficient": False,
                "passes_rvc": False,
                "unmet_qty_total": 1.0,
                "non_origin_value_usd": 70.0,
                "materials": [
                    {
                        "original_material_code": "ORIG",
                        "custom_code_basis": "FAMILY",
                        "need_qty": 10.0,
                        "exact_allocated_qty": 0.0,
                        "replacement_allocated_qty": 9.0,
                        "unmet_qty": 1.0,
                        "exact_allocations": [],
                        "replacement_allocations": [
                            {
                                "bucket_id": "b1",
                                "tracking_key": "t1",
                                "material_code": "REPL",
                                "custom_code_basis": "FAMILY",
                                "declaration_no": "D1",
                                "declaration_item_no": 1,
                                "import_date": "2026-04-01",
                                "allocated_qty": 9.0,
                                "unit_price_usd": 7.0,
                                "exchange_rate": 1.0,
                                "origin": "CHINA",
                                "source": "BCCT",
                                "source_row_no": "1",
                                "candidate_class": "confirmed_clean",
                                "mapping_status": "confirmed_declared_equals_dm_code",
                                "mapping_confidence": "medium",
                                "evidence_source": "mapping",
                                "staff_known_substitute": False,
                            }
                        ],
                    }
                ],
            },
            "material_mappings": [],
            "local_stock_state_after_product": {},
        }
        export_line = FakeExportLine(
            model_code="MODEL-A",
            declaration_no="DECL",
            declaration_item_no=1,
            export_date=date(2026, 4, 10),
            quantity=10.0,
            unit_price_usd=10.0,
        )
        setattr(export_line, "name", "Model A")
        setattr(export_line, "hs_code", "85044090")

        with tempfile.TemporaryDirectory() as tmpdir:
            tmp = Path(tmpdir)
            rvc_path = tmp / "replacement-excel" / "seed" / "replacement-rvc-standard.xlsx"
            bom_path = tmp / "replacement-excel" / "seed" / "replacement-bom-before-after.xlsx"
            runner.write_rvc_standard_workbook(
                rvc_path,
                [snapshot],
                {("MODEL-A", "DECL"): export_line},
                {"REPL": {"name": "Replacement Material", "hs_code": "850440", "unit": "PIECES"}},
            )
            runner.write_bom_before_after_workbook(bom_path, [snapshot])

            rvc_wb = openpyxl.load_workbook(rvc_path, data_only=True)
            bom_wb = openpyxl.load_workbook(bom_path, data_only=True)

        self.assertIn("Thiếu NVL", rvc_wb["01-MODEL-A"]["M13"].value)
        self.assertEqual(rvc_wb["01-MODEL-A"]["A13"].fill.fill_type, "solid")
        self.assertEqual(bom_wb["01-MODEL-A"]["A8"].fill.fill_type, "solid")
        self.assertEqual(bom_wb["01-MODEL-A"]["I8"].fill.fill_type, "solid")

    def test_write_stock_tracking_workbook_creates_expected_sheets(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            workbook_path = Path(tmpdir) / "replacement-excel" / "seed" / "replacement-stock-tracking.xlsx"
            runner.write_stock_tracking_workbook(
                workbook_path,
                initial_bucket_rows=[
                    {
                        "material_code": "ORIG",
                        "bucket_material_code": "ORIG",
                        "candidate_material_code": "ORIG",
                        "custom_code_basis": "FAMILY",
                        "bucket_id": "b1",
                        "tracking_key": "t1",
                        "declaration_no": "D1",
                        "declaration_item_no": 1,
                        "import_date": date(2026, 4, 1),
                        "source": "BCCT",
                        "source_row_no": "1",
                        "hs_code": "HS",
                        "name": "Material",
                        "origin": "CHINA",
                        "unit_price_usd": 1.0,
                        "exchange_rate": 1.0,
                        "remaining_qty": 10.0,
                        "candidate_class": "confirmed_clean",
                        "mapping_status": "confirmed_declared_equals_dm_code",
                        "mapping_confidence": "medium",
                        "evidence_source": "mapping",
                        "shipment_variant_scope": "confirmed",
                        "variant_scope": "",
                        "variant_hit_count": 1,
                    }
                ],
                consumption_rows=[
                    {
                        "starting_point_id": "seed",
                        "sequence_no": 1,
                        "model_code": "MODEL-A",
                        "variant_id": "MODEL-A__block1",
                        "bom_code": "MODEL-A",
                        "declaration_no": "DECL",
                        "original_material_code": "ORIG",
                        "allocation_type": "replacement",
                        "allocated_material_code": "REPL",
                        "custom_code_basis": "FAMILY",
                        "bucket_id": "b1",
                        "tracking_key": "t1",
                        "bucket_declaration_no": "D1",
                        "bucket_declaration_item_no": 1,
                        "import_date": date(2026, 4, 1),
                        "allocated_qty": 5.0,
                        "unit_price_usd": 1.0,
                        "allocated_value_usd": 5.0,
                        "origin": "CHINA",
                        "source": "BCCT",
                        "source_row_no": "1",
                        "candidate_class": "confirmed_clean",
                        "mapping_status": "confirmed_declared_equals_dm_code",
                        "mapping_confidence": "medium",
                        "evidence_source": "mapping",
                        "staff_known_substitute": False,
                    }
                ],
                after_product_rows=[
                    {
                        "starting_point_id": "seed",
                        "sequence_no": 1,
                        "model_code": "MODEL-A",
                        "declaration_no": "DECL",
                        "material_code": "ORIG",
                        "remaining_qty": 5.0,
                        "bucket_count": 1,
                    }
                ],
                final_bucket_rows=[
                    {
                        "material_code": "ORIG",
                        "bucket_material_code": "ORIG",
                        "candidate_material_code": "ORIG",
                        "custom_code_basis": "FAMILY",
                        "bucket_id": "b1",
                        "tracking_key": "t1",
                        "declaration_no": "D1",
                        "declaration_item_no": 1,
                        "import_date": date(2026, 4, 1),
                        "source": "BCCT",
                        "source_row_no": "1",
                        "hs_code": "HS",
                        "name": "Material",
                        "origin": "CHINA",
                        "unit_price_usd": 1.0,
                        "exchange_rate": 1.0,
                        "remaining_qty": 5.0,
                        "candidate_class": "confirmed_clean",
                        "mapping_status": "confirmed_declared_equals_dm_code",
                        "mapping_confidence": "medium",
                        "evidence_source": "mapping",
                        "shipment_variant_scope": "confirmed",
                        "variant_scope": "",
                        "variant_hit_count": 1,
                    }
                ],
                final_material_rows=[
                    {
                        "material_code": "ORIG",
                        "remaining_qty": 5.0,
                        "bucket_count": 1,
                    }
                ],
            )
            workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)

        self.assertEqual(
            workbook.sheetnames,
            ["Initial Buckets", "Consumption Ledger", "Stock After Product", "Final Buckets", "Final By Material"],
        )

    def test_build_candidate_rows_technical_reference_filters_to_whitelist(self) -> None:
        material = runner.MaterialState(
            original_material_code="ORIG",
            need_qty=10.0,
            exact_allocated_qty=0.0,
            unmet_qty=10.0,
            custom_code_basis="FAMILY",
        )
        stock_by_basis = {
            "FAMILY": [
                make_bucket(
                    basis="FAMILY",
                    material="REPL-ALLOWED",
                    bucket_id="allowed",
                    qty=10.0,
                    unit_price=1.0,
                    import_date=date(2026, 4, 1),
                ),
                make_bucket(
                    basis="FAMILY",
                    material="REPL-BLOCKED",
                    bucket_id="blocked",
                    qty=10.0,
                    unit_price=1.0,
                    import_date=date(2026, 4, 1),
                ),
            ]
        }
        stock_by_material = {
            "REPL-ALLOWED": [stock_by_basis["FAMILY"][0]],
            "REPL-BLOCKED": [stock_by_basis["FAMILY"][1]],
        }
        rows = runner.build_candidate_rows(
            self.baseline,
            runner.CandidateContext(
                starting_point_id="seed",
                sequence_no=1,
                iteration_no=1,
                model_code=self.export_line.model_code,
                variant_id=self.variant.variant_id,
                declaration_no=self.export_line.declaration_no,
            ),
            self.export_line,
            self.variant,
            [material],
            runner.summarize_product(self.baseline, self.export_line, [material]),
            material,
            stock_by_basis,
            set(),
            import_lead_days=2,
            max_import_age_days=0,
            stock_by_material=stock_by_material,
            replacement_mode=runner.REPLACEMENT_MODE_TECHNICAL_REFERENCE,
            explicit_substitute_index={("MODEL-A", "ORIG"): {"REPL-ALLOWED"}},
        )
        self.assertEqual([row["candidate_material_code"] for row in rows], ["REPL-ALLOWED"])

    def test_technical_reference_mode_can_use_cross_basis_explicit_substitute(self) -> None:
        initial_stock = {
            "ORIG": [
                make_bucket(
                    basis="FAMILY-ORIG",
                    material="ORIG",
                    bucket_id="orig-expensive",
                    qty=1.0,
                    unit_price=80.0,
                    import_date=date(2026, 4, 1),
                )
            ],
            "ALT-VN": [
                make_bucket(
                    basis="FAMILY-ALT",
                    material="ALT-VN",
                    bucket_id="alt-vn",
                    qty=1.0,
                    unit_price=5.0,
                    import_date=date(2026, 4, 1),
                    origin="VIETNAM",
                )
            ],
        }
        basis_stats = {
            "ORIG": {
                "FAMILY-ORIG": runner.BasisStats(
                    custom_code_basis="FAMILY-ORIG",
                    distinct_material_codes={"ORIG"},
                    total_qty=1.0,
                    confirmed_qty=1.0,
                )
            }
        }
        variant = FakeVariant("MODEL-A__block1", "MODEL-A", [FakeBomLine("ORIG", 1.0)])
        export_line = FakeExportLine("MODEL-A", "DECL-1", 1, date(2026, 4, 10), 1.0, 100.0)

        with tempfile.TemporaryDirectory() as tmpdir:
            tmp_path = Path(tmpdir)
            heuristic_result = runner.run_starting_point(
                self.baseline,
                shipment_id="SHIP",
                starting_point_id="seed",
                export_lines=[export_line],
                variant_by_id={"MODEL-A": variant},
                initial_stock_by_material=initial_stock,
                material_basis_stats=basis_stats,
                staff_substitutes=set(),
                import_lead_days=2,
                max_import_age_days=0,
                output_dir=tmp_path / "heuristic",
            )
            reference_result = runner.run_starting_point(
                self.baseline,
                shipment_id="SHIP",
                starting_point_id="seed",
                export_lines=[export_line],
                variant_by_id={"MODEL-A": variant},
                initial_stock_by_material=initial_stock,
                material_basis_stats=basis_stats,
                staff_substitutes=set(),
                import_lead_days=2,
                max_import_age_days=0,
                output_dir=tmp_path / "reference",
                replacement_mode=runner.REPLACEMENT_MODE_TECHNICAL_REFERENCE,
                explicit_substitute_index={("MODEL-A", "ORIG"): {"ALT-VN"}},
            )

        self.assertFalse(heuristic_result["seed_summary"]["product_status"][0]["passes_after"])
        self.assertTrue(reference_result["seed_summary"]["product_status"][0]["passes_after"])
        self.assertEqual(reference_result["seed_summary"]["product_status"][0]["changed_materials"], "ORIG")


if __name__ == "__main__":
    unittest.main()
