#!/usr/bin/env python3

from __future__ import annotations

import argparse
import importlib.util
import sys
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font


DEFAULT_ROOT_WORKBOOK = Path(
    "data/extracted/Growatt-20260421/Growatt/"
    "tru lui CO final  SXXK - 2025 commercial-MAC - Huyền đúng.xlsm"
)
DEFAULT_UNPACKED_WORKBOOK = Path(
    "data/extracted/Growatt-20260421/unpacked/lo-da-lam/"
    "CO cần chỉnh bảng kê tỷ lệ NVL (lô đã làm)/"
    "tru lui CO final  SXXK - 2025 commercial-MAC - Huyền đúng.xlsm"
)
DEFAULT_OUTPUT_DIR = Path("data/cases/growatt-rvc-20260421/shared/normalized")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root-workbook", type=Path, default=DEFAULT_ROOT_WORKBOOK)
    parser.add_argument("--unpacked-workbook", type=Path, default=DEFAULT_UNPACKED_WORKBOOK)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    return parser.parse_args()


def load_case_data_module() -> object:
    module_path = Path(__file__).with_name("build-growatt-case-data.py")
    spec = importlib.util.spec_from_file_location("growatt_case_data_runtime", module_path)
    if spec is None or spec.loader is None:
        raise SystemExit(f"Unable to load {module_path}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def build_multi_block_rows(dm_rows: list[dict[str, object]]) -> list[dict[str, object]]:
    by_bom: dict[str, list[dict[str, object]]] = defaultdict(list)
    for row in dm_rows:
        by_bom[str(row["bom_code"])].append(row)

    output_rows: list[dict[str, object]] = []
    for bom_code, rows in sorted(by_bom.items()):
        variants = sorted({str(item["bom_variant_id"]) for item in rows})
        if len(variants) <= 1:
            continue
        for variant_id in variants:
            variant_rows = [item for item in rows if str(item["bom_variant_id"]) == variant_id]
            output_rows.append(
                {
                    "bom_code": bom_code,
                    "variant_count": len(variants),
                    "variant_id": variant_id,
                    "row_start": int(variant_rows[0]["dm_row_no"]),
                    "row_end": int(variant_rows[-1]["dm_row_no"]),
                    "line_count": len(variant_rows),
                    "first_ordinal_key": variant_rows[0]["ordinal_key"],
                    "last_ordinal_key": variant_rows[-1]["ordinal_key"],
                }
            )
    return output_rows


def build_duplicate_material_rows(dm_rows: list[dict[str, object]]) -> list[dict[str, object]]:
    by_variant: dict[str, list[dict[str, object]]] = defaultdict(list)
    for row in dm_rows:
        by_variant[str(row["bom_variant_id"])].append(row)

    output_rows: list[dict[str, object]] = []
    for variant_id, rows in sorted(by_variant.items()):
        bom_code = str(rows[0]["bom_code"])
        counts = Counter(str(item["material_code"]) for item in rows)
        for material_code, occurrences in sorted(counts.items()):
            if occurrences <= 1:
                continue
            dup_rows = [item for item in rows if str(item["material_code"]) == material_code]
            output_rows.append(
                {
                    "bom_code": bom_code,
                    "variant_id": variant_id,
                    "material_code": material_code,
                    "occurrences": occurrences,
                    "total_qty_per_unit": sum(float(item["qty_per_unit"]) for item in dup_rows),
                    "row_start": int(dup_rows[0]["dm_row_no"]),
                    "row_end": int(dup_rows[-1]["dm_row_no"]),
                    "dm_row_nos": ", ".join(str(int(item["dm_row_no"])) for item in dup_rows),
                    "ordinal_keys": ", ".join(str(item["ordinal_key"]) for item in dup_rows),
                    "qty_per_unit_rows": ", ".join(str(item["qty_per_unit"]) for item in dup_rows),
                }
            )
    return output_rows


def write_sheet(
    workbook: openpyxl.Workbook,
    title: str,
    rows: list[dict[str, object]],
    fieldnames: list[str],
) -> None:
    sheet = workbook.create_sheet(title[:31])
    sheet.append(fieldnames)
    for row in rows:
        sheet.append([row.get(field) for field in fieldnames])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def write_audit_workbook(
    path: Path,
    source_workbook: Path,
    multi_block_rows: list[dict[str, object]],
    duplicate_rows: list[dict[str, object]],
) -> None:
    workbook = openpyxl.Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary["A1"] = "File nguồn"
    summary["B1"] = str(source_workbook)
    summary["A2"] = "Số mã SP có nhiều block BOM"
    summary["B2"] = len({row["bom_code"] for row in multi_block_rows})
    summary["A3"] = "Số dòng trong sheet Multi Block BOMs"
    summary["B3"] = len(multi_block_rows)
    summary["A4"] = "Số case NVL bị trùng trong cùng 1 block BOM"
    summary["B4"] = len(duplicate_rows)
    summary["A6"] = "Giải thích"
    summary["A7"] = (
        "Sheet 'Multi Block BOMs' liệt kê các mã sản phẩm xuất hiện thành nhiều block BOM "
        "trong sheet DM, kèm khoảng dòng từ dòng nào đến dòng nào."
    )
    summary["A8"] = (
        "Sheet 'Duplicate Materials' liệt kê các mã NVL bị lặp trong cùng 1 block BOM. "
        "Mỗi dòng là 1 case dạng (mã SP, block BOM, mã NVL), không phải 1 dòng BOM gốc."
    )
    summary["A9"] = (
        "Ví dụ nếu một mã NVL xuất hiện 4 lần trong cùng 1 block BOM thì sheet này ghi 1 dòng "
        "với occurrences = 4, đồng thời liệt kê các dm_row_no và ordinal_key tương ứng."
    )
    summary["A10"] = (
        "Sheet này chưa kết luận các dòng trùng là đúng hay sai nghiệp vụ; mục đích là để staff "
        "xác nhận nên giữ riêng từng dòng hay gộp theo mã NVL."
    )
    for cell in ("A1", "A2", "A3", "A4", "A6"):
        summary[cell].font = Font(bold=True)
    for row_idx in range(1, 11):
        summary[f"A{row_idx}"].alignment = Alignment(vertical="top", wrap_text=True)
        summary[f"B{row_idx}"].alignment = Alignment(vertical="top", wrap_text=True)
    summary.column_dimensions["A"].width = 42
    summary.column_dimensions["B"].width = 120

    write_sheet(
        workbook,
        "Multi Block BOMs",
        multi_block_rows,
        [
            "bom_code",
            "variant_count",
            "variant_id",
            "row_start",
            "row_end",
            "line_count",
            "first_ordinal_key",
            "last_ordinal_key",
        ],
    )
    write_sheet(
        workbook,
        "Duplicate Materials",
        duplicate_rows,
        [
            "bom_code",
            "variant_id",
            "material_code",
            "occurrences",
            "total_qty_per_unit",
            "row_start",
            "row_end",
            "dm_row_nos",
            "ordinal_keys",
            "qty_per_unit_rows",
        ],
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def run(args: argparse.Namespace) -> list[Path]:
    case_data = load_case_data_module()
    workbooks = [
        ("growatt-root", args.root_workbook),
        ("growatt-lo-da-lam", args.unpacked_workbook),
    ]
    outputs: list[Path] = []
    for slug, workbook_path in workbooks:
        dm_rows = case_data.load_dm_variants(workbook_path)
        multi_block_rows = build_multi_block_rows(dm_rows)
        duplicate_rows = build_duplicate_material_rows(dm_rows)
        output_path = args.output_dir / f"dm-audit-{slug}.xlsx"
        write_audit_workbook(output_path, workbook_path, multi_block_rows, duplicate_rows)
        outputs.append(output_path)
    return outputs


def main() -> None:
    outputs = run(parse_args())
    for path in outputs:
        print(path)


if __name__ == "__main__":
    main()
