#!/usr/bin/env python3

from __future__ import annotations

import argparse
import csv
import json
import re
import xml.etree.ElementTree as ET
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from zipfile import ZipFile

import openpyxl
import xlrd


NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
DEFAULT_WORKBOOK = Path(
    "data/extracted/Growatt-20260421/Growatt/"
    "tru lui CO final  SXXK - 2025 commercial-MAC - Huyền đúng.xlsm"
)
DEFAULT_NK_REPORT = Path(
    "data/extracted/Growatt-20260421/Growatt/BaoCaoHangChiTietnk grw t3-t4.2026.xls"
)
DEFAULT_XK_REPORT = Path(
    "data/extracted/Growatt-20260421/Growatt/BaoCaoHangChiTietXK GRW T3-T4.2026 moi.xls"
)
DEFAULT_CASE_DIR = Path("data/cases/growatt-rvc-20260421/shared")
DEFAULT_USD_FX_WORKBOOK = Path("data/reference/DS_ty_gia_ngoai_te.xlsx")
ERP_CODE_RE = re.compile(r"^[A-Z0-9]+(?:\.[A-Z0-9]+)+(?:-[A-Z0-9-]+)?$")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--nk-report", type=Path, default=DEFAULT_NK_REPORT)
    parser.add_argument("--xk-report", type=Path, default=DEFAULT_XK_REPORT)
    parser.add_argument("--case-dir", type=Path, default=DEFAULT_CASE_DIR)
    parser.add_argument("--usd-fx-workbook", type=Path, default=DEFAULT_USD_FX_WORKBOOK)
    return parser.parse_args()


def write_csv(path: Path, rows: list[dict[str, object]], fieldnames: list[str]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows([{field: row.get(field) for field in fieldnames} for row in rows])


def clean_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def normalize_code_text(value: object, field_name: str) -> str:
    if value in (None, ""):
        return ""
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, int):
        raise ValueError(f"{field_name} must be stored as text, got integer {value!r}")
    if isinstance(value, float):
        raise ValueError(f"{field_name} must be stored as text, got numeric {value!r}")
    return str(value).strip()


def normalize_integral_identifier(value: object, field_name: str) -> str:
    if value in (None, ""):
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        raise ValueError(f"{field_name} must be an integer-like identifier, got {value!r}")
    text = clean_text(value)
    if re.fullmatch(r"\d+\.0+", text):
        return text.split(".", 1)[0]
    return text


def parse_float(value: object) -> float:
    text = clean_text(value)
    if not text:
        return 0.0
    return float(text)


def make_line_key(declaration_no: object, declaration_item_no: object) -> str:
    return f"{clean_text(declaration_no)}-{clean_text(declaration_item_no)}"


def to_date(value: object, datemode: int | None = None) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and datemode is not None:
        try:
            return datetime(*xlrd.xldate_as_tuple(value, datemode)).date()
        except Exception:
            return None
    if isinstance(value, str):
        text = value.strip()
        for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%m/%d/%Y"):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                continue
    return None


def looks_like_erp_code(text: str) -> bool:
    return bool(text and ERP_CODE_RE.match(text))


def parse_vnd_rate_text(value: object) -> float:
    digits = re.sub(r"[^0-9]", "", clean_text(value))
    if not digits:
        return 0.0
    return float(digits)


def load_usd_customs_rates(path: Path) -> list[tuple[date, float]]:
    if not path.exists():
        return []
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows: list[tuple[date, float]] = []
    for row in sheet.iter_rows(min_row=3, values_only=True):
        if clean_text(row[0]).upper() != "USD":
            continue
        effective_date = to_date(row[2])
        rate_value = parse_vnd_rate_text(row[3])
        if effective_date is None or rate_value <= 0:
            continue
        rows.append((effective_date, rate_value))
    rows.sort(key=lambda item: item[0])
    return rows


def lookup_usd_customs_rate(declaration_date: date | None, usd_customs_rates: list[tuple[date, float]]) -> float | None:
    if declaration_date is None:
        return None
    applicable_rate: float | None = None
    for effective_date, rate_value in usd_customs_rates:
        if effective_date <= declaration_date:
            applicable_rate = rate_value
            continue
        break
    return applicable_rate


def normalize_price_fields(
    raw_unit_price: object,
    tax_unit_price: object,
    source_exchange_rate: object,
    declaration_date: date | None,
    usd_customs_rates: list[tuple[date, float]],
) -> dict[str, object]:
    raw_value = parse_float(raw_unit_price)
    tax_value = parse_float(tax_unit_price)
    source_rate = parse_float(source_exchange_rate)
    customs_usd_rate = lookup_usd_customs_rate(declaration_date, usd_customs_rates)

    if tax_value > 0 and customs_usd_rate:
        return {
            "source_unit_price": raw_value,
            "unit_price_usd": tax_value / customs_usd_rate,
            "tax_unit_price": tax_value,
            "exchange_rate": customs_usd_rate,
            "source_exchange_rate": source_rate,
            "price_normalization_basis": "tax_vnd_over_customs_usd_rate",
        }
    if tax_value > 0 and source_rate > 1:
        return {
            "source_unit_price": raw_value,
            "unit_price_usd": tax_value / source_rate,
            "tax_unit_price": tax_value,
            "exchange_rate": source_rate,
            "source_exchange_rate": source_rate,
            "price_normalization_basis": "tax_vnd_over_source_exchange_rate",
        }
    if raw_value > 0:
        return {
            "source_unit_price": raw_value,
            "unit_price_usd": raw_value,
            "tax_unit_price": tax_value,
            "exchange_rate": customs_usd_rate or source_rate,
            "source_exchange_rate": source_rate,
            "price_normalization_basis": "raw_unit_price_field",
        }
    return {
        "source_unit_price": raw_value,
        "unit_price_usd": 0.0,
        "tax_unit_price": tax_value,
        "exchange_rate": customs_usd_rate or source_rate,
        "source_exchange_rate": source_rate,
        "price_normalization_basis": "missing_price",
    }


def split_name_parts(name: str) -> tuple[str, str]:
    if "#&" not in name:
        return "", name.strip()
    left, right = name.split("#&", 1)
    return left.strip(), right.strip()


def extract_paren_codes(name: str) -> list[str]:
    codes: list[str] = []
    for item in re.findall(r"\(([^)]+)\)", name):
        token = item.strip()
        if looks_like_erp_code(token):
            codes.append(token)
    return codes


def derive_name_fields(name: str, declared_code: str) -> dict[str, object]:
    label_code, description = split_name_parts(name)
    paren_codes = extract_paren_codes(name)
    if paren_codes:
        candidate_lookup_code = paren_codes[-1]
        extraction_status = "extracted_from_paren"
    elif looks_like_erp_code(declared_code):
        candidate_lookup_code = declared_code
        extraction_status = "fallback_declared_code"
    else:
        candidate_lookup_code = ""
        extraction_status = "missing_code"
    return {
        "label_code": label_code,
        "description_clean": description,
        "paren_code_candidates": ";".join(paren_codes),
        "candidate_lookup_code": candidate_lookup_code,
        "code_extraction_status": extraction_status,
    }


def annotate_lookup_resolution(
    rows: list[dict[str, object]],
    candidate_field: str,
    dm_codes: set[str],
    dm_family_codes: set[str],
    row_kind: str,
) -> None:
    for row in rows:
        candidate_code = clean_text(row.get(candidate_field))
        declared_code = clean_text(row.get("declared_code"))
        row["confirmed_lookup_code"] = ""
        row["final_lookup_key"] = ""
        if not candidate_code:
            row["dm_match_status"] = "no_lookup_code"
            row["mapping_status"] = "no_lookup_code"
            row["lookup_confidence"] = "none"
        elif candidate_code in dm_codes:
            row["dm_match_status"] = "exact_dm_match"
            if row_kind == "export":
                row["mapping_status"] = "confirmed_export_exact_dm_match"
                row["lookup_confidence"] = "high"
                row["confirmed_lookup_code"] = candidate_code
                row["final_lookup_key"] = candidate_code
            elif declared_code and declared_code == candidate_code:
                row["mapping_status"] = "confirmed_declared_equals_dm_code"
                row["lookup_confidence"] = "medium"
                row["confirmed_lookup_code"] = candidate_code
                row["final_lookup_key"] = candidate_code
            else:
                row["mapping_status"] = "candidate_exact_dm_match"
                row["lookup_confidence"] = "candidate"
        elif candidate_code.split("-", 1)[0] in dm_family_codes:
            row["dm_match_status"] = "family_dm_match"
            row["mapping_status"] = "candidate_family_dm_match"
            row["lookup_confidence"] = "candidate"
        else:
            row["dm_match_status"] = "not_in_dm"
            row["mapping_status"] = "candidate_not_in_dm"
            row["lookup_confidence"] = "low"


def reconcile_import_rows_against_nk2(
    bcct_rows: list[dict[str, object]],
    nk2_rows: list[dict[str, object]],
) -> None:
    confirmed_nk2_by_key: dict[str, dict[str, object]] = {}
    ambiguous_keys: set[str] = set()
    for row in nk2_rows:
        key = clean_text(row.get("tracking_key"))
        confirmed_code = clean_text(row.get("confirmed_lookup_code"))
        if not key or not confirmed_code:
            continue
        existing = confirmed_nk2_by_key.get(key)
        if existing is None:
            confirmed_nk2_by_key[key] = row
            continue
        if clean_text(existing.get("confirmed_lookup_code")) != confirmed_code:
            ambiguous_keys.add(key)

    for key in ambiguous_keys:
        confirmed_nk2_by_key.pop(key, None)

    for row in bcct_rows:
        row["matched_nk2_source_row_no"] = ""
        if clean_text(row.get("confirmed_lookup_code")):
            continue
        nk2_row = confirmed_nk2_by_key.get(clean_text(row.get("tracking_key")))
        if nk2_row is None:
            continue
        confirmed_code = clean_text(nk2_row.get("confirmed_lookup_code"))
        candidate_code = clean_text(row.get("lookup_material_code"))
        if candidate_code and candidate_code != confirmed_code:
            continue
        row["confirmed_lookup_code"] = confirmed_code
        row["final_lookup_key"] = confirmed_code
        row["mapping_status"] = "confirmed_via_nk2_same_customs_line"
        row["lookup_confidence"] = "medium"
        row["matched_nk2_source_row_no"] = nk2_row.get("source_row_no", "")


def load_shared_strings(archive: ZipFile) -> list[str]:
    root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
    strings: list[str] = []
    for si in root.findall(f"{NS}si"):
        strings.append("".join(text.text or "" for text in si.iter(f"{NS}t")))
    return strings


def cell_value(cell: ET.Element, shared_strings: list[str]) -> str | None:
    if cell.attrib.get("t") == "inlineStr":
        return "".join(text.text or "" for text in cell.iter(f"{NS}t"))
    raw = cell.find(f"{NS}v")
    if raw is None:
        return None
    value = raw.text or ""
    if cell.attrib.get("t") == "s":
        return shared_strings[int(value)]
    return value


def load_dm_variants(workbook_path: Path) -> list[dict[str, object]]:
    rows_out: list[dict[str, object]] = []
    with ZipFile(workbook_path) as archive:
        shared_strings = load_shared_strings(archive)
        rows: list[tuple[int, str, str, str, float]] = []
        for _, elem in ET.iterparse(archive.open("xl/worksheets/sheet4.xml"), events=("end",)):
            if elem.tag != f"{NS}row":
                continue
            row_no = int(elem.attrib["r"])
            if row_no < 7:
                elem.clear()
                continue
            values: dict[str, str | None] = {}
            for cell in elem.findall(f"{NS}c"):
                ref = cell.attrib["r"]
                col = "".join(ch for ch in ref if ch.isalpha())
                if col in {"A", "F"} and cell.attrib.get("t") not in {"s", "str", "inlineStr"}:
                    raise ValueError(f"DM code cell {ref} must be stored as text")
                values[col] = cell_value(cell, shared_strings)
            bom_code = normalize_code_text(values.get("A"), f"DM bom_code row {row_no}")
            ordinal_key = clean_text(values.get("B"))
            material_code = normalize_code_text(values.get("F"), f"DM material_code row {row_no}")
            qty_raw = values.get("I")
            if bom_code and material_code and qty_raw not in (None, ""):
                rows.append((row_no, bom_code, ordinal_key, material_code, float(qty_raw)))
            elem.clear()

    block_counts: dict[str, int] = defaultdict(int)
    current_code = ""
    current_rows: list[tuple[int, str, str, float]] = []

    def flush_current() -> None:
        nonlocal current_code, current_rows
        if not current_code or not current_rows:
            current_code = ""
            current_rows = []
            return
        block_counts[current_code] += 1
        variant_id = f"{current_code}__block{block_counts[current_code]}"
        family_code = current_code.split("-", 1)[0]
        for row_no, ordinal_key, material_code, qty_per_unit in current_rows:
            rows_out.append(
                {
                    "bom_code": current_code,
                    "product_family_code": family_code,
                    "bom_variant_id": variant_id,
                    "dm_row_no": row_no,
                    "ordinal_key": ordinal_key,
                    "material_code": material_code,
                    "qty_per_unit": qty_per_unit,
                }
            )
        current_code = ""
        current_rows = []

    for row_no, bom_code, ordinal_key, material_code, qty_per_unit in rows:
        if bom_code != current_code:
            flush_current()
            current_code = bom_code
        current_rows.append((row_no, ordinal_key, material_code, qty_per_unit))
    flush_current()
    return rows_out


def load_xk_rows(report_path: Path) -> list[dict[str, object]]:
    book = xlrd.open_workbook(str(report_path))
    sheet = book.sheet_by_index(0)
    rows: list[dict[str, object]] = []
    for row_idx in range(10, sheet.nrows):
        row = sheet.row_values(row_idx)
        name = clean_text(row[22])
        if not name:
            continue
        declared_code = normalize_code_text(row[20], f"XK declared_code row {row_idx + 1}")
        name_fields = derive_name_fields(name, declared_code)
        declaration_no = normalize_integral_identifier(row[1], f"XK declaration_no row {row_idx + 1}")
        declaration_item_no = normalize_integral_identifier(row[19], f"XK declaration_item_no row {row_idx + 1}")
        rows.append(
            {
                "source_row_no": row_idx + 1,
                "declaration_no": declaration_no,
                "declaration_date": to_date(row[2], book.datemode),
                "declaration_item_no": declaration_item_no,
                "shipment_id": clean_text(row[50]),
                "declared_code": declared_code,
                "hs_code": clean_text(row[21]),
                "name": name,
                "origin": clean_text(row[23]),
                "unit_price_usd": float(row[24] or 0),
                "quantity": float(row[26] or 0),
                "unit": clean_text(row[27]),
                "invoice_date": to_date(row[51], book.datemode),
                "invoice_no": clean_text(row[50]),
                "incoterm": clean_text(row[17]),
                "internal_code_for_dm": name_fields["candidate_lookup_code"],
                **name_fields,
            }
        )
    return rows


def load_bcct_nk_rows(report_path: Path, usd_customs_rates: list[tuple[date, float]]) -> list[dict[str, object]]:
    book = xlrd.open_workbook(str(report_path))
    sheet = book.sheet_by_index(0)
    rows: list[dict[str, object]] = []
    for row_idx in range(10, sheet.nrows):
        row = sheet.row_values(row_idx)
        name = clean_text(row[22])
        declared_code = normalize_code_text(row[20], f"NK report declared_code row {row_idx + 1}")
        if not declared_code and not name:
            continue
        name_fields = derive_name_fields(name, declared_code)
        declaration_date = to_date(row[2], book.datemode)
        price_fields = normalize_price_fields(row[24], row[25], row[10], declaration_date, usd_customs_rates)
        declaration_no = normalize_integral_identifier(row[1], f"NK report declaration_no row {row_idx + 1}")
        declaration_item_no = normalize_integral_identifier(row[19], f"NK report declaration_item_no row {row_idx + 1}")
        rows.append(
            {
                "source": "BCCT_NK",
                "source_row_no": row_idx + 1,
                "declaration_no": declaration_no,
                "declaration_date": declaration_date,
                "declaration_item_no": declaration_item_no,
                "tracking_key": make_line_key(declaration_no, declaration_item_no),
                "declared_code": declared_code,
                "lookup_material_code": name_fields["candidate_lookup_code"],
                "hs_code": clean_text(row[21]),
                "name": name,
                "origin": clean_text(row[23]),
                "import_qty": float(row[26] or 0),
                "unit": clean_text(row[27]),
                "partner_name": clean_text(row[49]),
                "invoice_no": clean_text(row[50]),
                "invoice_date": to_date(row[51], book.datemode),
                **price_fields,
                **name_fields,
            }
        )
    return rows


def load_nk2_rows(workbook_path: Path, usd_customs_rates: list[tuple[date, float]]) -> list[dict[str, object]]:
    workbook = openpyxl.load_workbook(
        workbook_path,
        data_only=True,
        read_only=True,
        keep_vba=True,
    )
    sheet = workbook["NK2"]
    rows: list[dict[str, object]] = []
    for row_no, row in enumerate(sheet.iter_rows(min_row=5, values_only=True), start=5):
        declared_code = normalize_code_text(row[4], f"NK2 declared_code row {row_no}")
        name = clean_text(row[6])
        if not declared_code and not name:
            continue
        name_fields = derive_name_fields(name, declared_code)
        declaration_date = to_date(row[1])
        price_fields = normalize_price_fields(row[8], row[9], row[15], declaration_date, usd_customs_rates)
        declaration_no = normalize_integral_identifier(row[0], f"NK2 declaration_no row {row_no}")
        declaration_item_no = normalize_integral_identifier(row[3], f"NK2 declaration_item_no row {row_no}")
        rows.append(
            {
                "source": "NK2",
                "source_row_no": row_no,
                "declaration_no": declaration_no,
                "declaration_date": declaration_date,
                "declaration_item_no": declaration_item_no,
                "tracking_key": make_line_key(declaration_no, declaration_item_no),
                "declared_code": declared_code,
                "lookup_material_code": name_fields["candidate_lookup_code"],
                "hs_code": clean_text(row[5]),
                "name": name,
                "origin": clean_text(row[7]),
                "import_qty": float(row[10] or 0),
                "unit": clean_text(row[11]),
                "partner_name": clean_text(row[12]),
                "invoice_no": clean_text(row[13]),
                "invoice_date": to_date(row[14]),
                "used_qty": float(row[16] or 0),
                "remaining_qty": float(row[17] or 0),
                **price_fields,
                **name_fields,
            }
        )
    return rows


def build_updated_stock_tracking(
    nk2_rows: list[dict[str, object]],
    bcct_rows: list[dict[str, object]],
) -> tuple[list[dict[str, object]], dict[str, int]]:
    nk2_by_key = {row["tracking_key"]: row for row in nk2_rows}
    bcct_by_key = {row["tracking_key"]: row for row in bcct_rows}
    all_keys = sorted(set(nk2_by_key) | set(bcct_by_key))
    rows_out: list[dict[str, object]] = []
    stats = {
        "nk2_rows": len(nk2_rows),
        "bcct_rows": len(bcct_rows),
        "merged_overlap_rows": 0,
        "nk2_only_rows": 0,
        "bcct_appended_rows": 0,
    }

    for key in all_keys:
        nk2_row = nk2_by_key.get(key)
        bcct_row = bcct_by_key.get(key)
        if nk2_row and bcct_row:
            row = {**bcct_row}
            row["source"] = "MERGED_NK2_BCCT"
            row["used_qty"] = nk2_row["used_qty"]
            row["remaining_qty"] = nk2_row["remaining_qty"]
            row["nk2_source_row_no"] = nk2_row["source_row_no"]
            row["bcct_source_row_no"] = bcct_row["source_row_no"]
            stats["merged_overlap_rows"] += 1
        elif nk2_row:
            row = dict(nk2_row)
            row["source"] = "NK2_ONLY"
            row["nk2_source_row_no"] = nk2_row["source_row_no"]
            row["bcct_source_row_no"] = None
            stats["nk2_only_rows"] += 1
        else:
            row = dict(bcct_row)
            row["source"] = "BCCT_APPENDED"
            row["used_qty"] = 0.0
            row["remaining_qty"] = row["import_qty"]
            row["nk2_source_row_no"] = None
            row["bcct_source_row_no"] = bcct_row["source_row_no"]
            stats["bcct_appended_rows"] += 1
        rows_out.append(row)

    return rows_out, stats


def write_xlsx(
    path: Path,
    updated_stock_rows: list[dict[str, object]],
    import_rows: list[dict[str, object]],
    export_rows: list[dict[str, object]],
    dm_rows: list[dict[str, object]],
    stats: dict[str, object],
) -> None:
    wb = openpyxl.Workbook()
    default = wb.active
    wb.remove(default)

    def add_sheet(name: str, rows: list[dict[str, object]]) -> None:
        ws = wb.create_sheet(name)
        if not rows:
            ws.append(["empty"])
            return
        fieldnames = list(rows[0].keys())
        text_fields = {
            "tracking_key",
            "declaration_no",
            "declaration_item_no",
            "declared_code",
            "lookup_material_code",
            "confirmed_lookup_code",
            "final_lookup_key",
            "matched_nk2_source_row_no",
            "hs_code",
            "label_code",
            "paren_code_candidates",
            "invoice_no",
            "nk2_source_row_no",
            "bcct_source_row_no",
            "shipment_id",
            "internal_code_for_dm",
            "incoterm",
            "bom_code",
            "product_family_code",
            "bom_variant_id",
            "material_code",
            "ordinal_key",
        }
        ws.append(fieldnames)
        for row in rows:
            ws.append([row.get(field) for field in fieldnames])
        for idx, field in enumerate(fieldnames, start=1):
            if field in text_fields:
                for row_idx in range(2, ws.max_row + 1):
                    ws.cell(row=row_idx, column=idx).number_format = "@"
        ws.freeze_panes = "A2"

    summary = wb.create_sheet("SUMMARY")
    for idx, (key, value) in enumerate(stats.items(), start=1):
        summary.cell(idx, 1).value = key
        summary.cell(idx, 2).value = json.dumps(value, ensure_ascii=False) if isinstance(value, dict) else value

    add_sheet("CO_STOCK_UPDATED", updated_stock_rows)
    add_sheet("IMPORTS_NORMALIZED", import_rows)
    add_sheet("EXPORTS_NORMALIZED", export_rows)
    add_sheet("DM_VARIANTS", dm_rows)
    wb.save(path)


def main() -> None:
    args = parse_args()
    normalized_dir = args.case_dir / "normalized"
    normalized_dir.mkdir(parents=True, exist_ok=True)

    dm_rows = load_dm_variants(args.workbook)
    dm_codes = {clean_text(row["bom_code"]) for row in dm_rows} | {clean_text(row["material_code"]) for row in dm_rows}
    dm_family_codes = {clean_text(row["product_family_code"]) for row in dm_rows}
    usd_customs_rates = load_usd_customs_rates(args.usd_fx_workbook)
    export_rows = load_xk_rows(args.xk_report)
    bcct_nk_rows = load_bcct_nk_rows(args.nk_report, usd_customs_rates)
    nk2_rows = load_nk2_rows(args.workbook, usd_customs_rates)
    annotate_lookup_resolution(
        export_rows,
        "internal_code_for_dm",
        dm_codes,
        dm_family_codes,
        row_kind="export",
    )
    annotate_lookup_resolution(
        bcct_nk_rows,
        "lookup_material_code",
        dm_codes,
        dm_family_codes,
        row_kind="import",
    )
    annotate_lookup_resolution(
        nk2_rows,
        "lookup_material_code",
        dm_codes,
        dm_family_codes,
        row_kind="stock",
    )
    reconcile_import_rows_against_nk2(bcct_nk_rows, nk2_rows)
    updated_stock_rows, stats = build_updated_stock_tracking(nk2_rows, bcct_nk_rows)

    import_extraction_counts = Counter(row["code_extraction_status"] for row in bcct_nk_rows)
    import_dm_match_counts = Counter(row["dm_match_status"] for row in bcct_nk_rows)
    import_mapping_counts = Counter(row["mapping_status"] for row in bcct_nk_rows)
    export_extraction_counts = Counter(row["code_extraction_status"] for row in export_rows)
    export_dm_match_counts = Counter(row["dm_match_status"] for row in export_rows)
    export_mapping_counts = Counter(row["mapping_status"] for row in export_rows)
    stock_mapping_counts = Counter(row["mapping_status"] for row in updated_stock_rows)
    import_price_basis_counts = Counter(row["price_normalization_basis"] for row in bcct_nk_rows)
    stock_price_basis_counts = Counter(row["price_normalization_basis"] for row in updated_stock_rows)

    manifest = {
        "generated_at": datetime.now().isoformat(),
        "workbook": str(args.workbook),
        "nk_report": str(args.nk_report),
        "xk_report": str(args.xk_report),
        "usd_fx_workbook": str(args.usd_fx_workbook) if args.usd_fx_workbook else "",
        "usd_fx_rate_rows": len(usd_customs_rates),
        **stats,
        "import_code_extraction_counts": dict(import_extraction_counts),
        "import_dm_match_counts": dict(import_dm_match_counts),
        "import_mapping_counts": dict(import_mapping_counts),
        "import_price_basis_counts": dict(import_price_basis_counts),
        "export_code_extraction_counts": dict(export_extraction_counts),
        "export_dm_match_counts": dict(export_dm_match_counts),
        "export_mapping_counts": dict(export_mapping_counts),
        "stock_mapping_counts": dict(stock_mapping_counts),
        "stock_price_basis_counts": dict(stock_price_basis_counts),
    }

    stock_fields = [
        "source",
        "tracking_key",
        "declaration_no",
        "declaration_date",
        "declaration_item_no",
        "declared_code",
        "lookup_material_code",
        "confirmed_lookup_code",
        "final_lookup_key",
        "matched_nk2_source_row_no",
        "code_extraction_status",
        "dm_match_status",
        "mapping_status",
        "lookup_confidence",
        "hs_code",
        "label_code",
        "paren_code_candidates",
        "name",
        "description_clean",
        "origin",
        "source_unit_price",
        "unit_price_usd",
        "tax_unit_price",
        "price_normalization_basis",
        "import_qty",
        "used_qty",
        "remaining_qty",
        "unit",
        "partner_name",
        "invoice_no",
        "invoice_date",
        "exchange_rate",
        "source_exchange_rate",
        "nk2_source_row_no",
        "bcct_source_row_no",
    ]
    import_fields = [
        "source",
        "source_row_no",
        "tracking_key",
        "declaration_no",
        "declaration_date",
        "declaration_item_no",
        "declared_code",
        "lookup_material_code",
        "confirmed_lookup_code",
        "final_lookup_key",
        "matched_nk2_source_row_no",
        "code_extraction_status",
        "dm_match_status",
        "mapping_status",
        "lookup_confidence",
        "hs_code",
        "label_code",
        "paren_code_candidates",
        "name",
        "description_clean",
        "origin",
        "source_unit_price",
        "unit_price_usd",
        "tax_unit_price",
        "price_normalization_basis",
        "import_qty",
        "unit",
        "partner_name",
        "invoice_no",
        "invoice_date",
        "exchange_rate",
        "source_exchange_rate",
    ]
    export_fields = [
        "source_row_no",
        "declaration_no",
        "declaration_date",
        "declaration_item_no",
        "shipment_id",
        "declared_code",
        "internal_code_for_dm",
        "confirmed_lookup_code",
        "final_lookup_key",
        "code_extraction_status",
        "dm_match_status",
        "mapping_status",
        "lookup_confidence",
        "hs_code",
        "label_code",
        "paren_code_candidates",
        "name",
        "description_clean",
        "origin",
        "unit_price_usd",
        "quantity",
        "unit",
        "invoice_no",
        "invoice_date",
        "incoterm",
    ]
    dm_fields = [
        "bom_code",
        "product_family_code",
        "bom_variant_id",
        "dm_row_no",
        "ordinal_key",
        "material_code",
        "qty_per_unit",
    ]

    write_csv(normalized_dir / "co-stock-tracking-updated.csv", updated_stock_rows, stock_fields)
    write_csv(normalized_dir / "imports-normalized.csv", bcct_nk_rows, import_fields)
    write_csv(normalized_dir / "exports-normalized.csv", export_rows, export_fields)
    write_csv(normalized_dir / "dm-variants.csv", dm_rows, dm_fields)

    (normalized_dir / "manifest.json").write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False, default=str),
        encoding="utf-8",
    )

    write_xlsx(
        normalized_dir / "growatt-case-data.xlsx",
        updated_stock_rows=updated_stock_rows,
        import_rows=bcct_nk_rows,
        export_rows=export_rows,
        dm_rows=dm_rows,
        stats=manifest,
    )

    print(f"Wrote {normalized_dir}")
    print(json.dumps(manifest, indent=2, ensure_ascii=False, default=str))


if __name__ == "__main__":
    main()
