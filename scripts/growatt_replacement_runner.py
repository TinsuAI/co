#!/usr/bin/env python3

from __future__ import annotations

import argparse
import copy
import csv
import importlib.util
import json
import re
import sys
from collections import defaultdict
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from growatt_case_config import resolve_shipment_policy
from growatt_bom_sources import technical_priority_artifact_dir


DEFAULT_CASE_DIR = Path("data/cases/growatt-rvc-20260421")
DEFAULT_SHIPMENT = "GIN01426B282"
DEFAULT_DOC = Path("docs/growatt-b282-replacement-runner.md")
EPS = 1e-9
COST_EPS = 1e-4
INTERNAL_MATERIAL_CODE_RE = re.compile(r"^[A-Z]*\d+(?:\.\d+)+(?:-[A-Z0-9-]+)?$")
REPLACEMENT_MODE_HEURISTIC_CUSTOM_BASIS = "heuristic_custom_basis"
REPLACEMENT_MODE_TECHNICAL_REFERENCE = "technical_reference_explicit"
STAFF_SUBSTITUTE_PATHS = (
    Path(
        "data/extracted/Growatt-20260421/unpacked/"
        "lo-da-lam/CO cần chỉnh bảng kê tỷ lệ NVL (lô đã làm)/NVL Thay thế.xlsx"
    ),
    Path(
        "data/extracted/Growatt-20260421/unpacked/"
        "lo-chua-lam-1/CO cần chỉnh bảng kê tỷ NVL (lô chưa làm)/NVL Thay thế.xlsx"
    ),
    Path(
        "data/extracted/Growatt-20260421/unpacked/"
        "lo-chua-lam-2/CO cần chỉnh bảng kê tỷ lệ NVL (lô chưa làm 2)/NVL Thay thế.xlsx"
    ),
)
MERCHANT_NAME = "CÔNG TY TNHH NĂNG LƯỢNG MỚI GROWATT VIỆT NAM"
MERCHANT_TAX_ID = "0202177200"
UNMET_FILL = PatternFill(fill_type="solid", fgColor="FCE4D6")
CHANGED_FILL = PatternFill(fill_type="solid", fgColor="FFF2CC")


@dataclass
class ReplacementBasisBucket:
    custom_code_basis: str
    erp_material_code: str
    bucket_id: str
    tracking_key: str
    declaration_no: str
    declaration_item_no: int
    import_date: date | None
    source: str
    source_row_no: str
    hs_code: str
    name: str
    origin: str
    unit_price_usd: float
    exchange_rate: float
    remaining_qty: float
    candidate_class: str
    mapping_status: str
    mapping_confidence: str
    evidence_source: str
    shipment_variant_scope: str
    variant_scope: tuple[str, ...]
    variant_hit_count: int
    candidate_material_code: str = ""


@dataclass
class BucketAllocation:
    bucket_id: str
    tracking_key: str
    material_code: str
    custom_code_basis: str
    declaration_no: str
    declaration_item_no: int
    import_date: date | None
    allocated_qty: float
    unit_price_usd: float
    exchange_rate: float
    origin: str
    source: str
    source_row_no: str
    candidate_class: str
    mapping_status: str
    mapping_confidence: str
    evidence_source: str
    staff_known_substitute: bool


@dataclass
class MaterialState:
    original_material_code: str
    need_qty: float
    exact_allocations: list[BucketAllocation] = field(default_factory=list)
    replacement_allocations: list[BucketAllocation] = field(default_factory=list)
    exact_allocated_qty: float = 0.0
    unmet_qty: float = 0.0
    custom_code_basis: str = ""

    @property
    def replacement_allocated_qty(self) -> float:
        return sum(item.allocated_qty for item in self.replacement_allocations)

    @property
    def total_allocated_qty(self) -> float:
        return self.exact_allocated_qty + self.replacement_allocated_qty

    @property
    def active_allocations(self) -> list[BucketAllocation]:
        return [*self.exact_allocations, *self.replacement_allocations]


@dataclass
class ProductSnapshot:
    rvc_percent: float | None
    margin_to_threshold: float | None
    stock_sufficient: bool
    passes_rvc: bool
    unmet_qty_total: float
    non_origin_value_usd: float
    materials: list[MaterialState]


@dataclass
class MaterialPlan:
    original_material_code: str
    custom_code_basis: str
    allocations: list[BucketAllocation]
    covered_qty: float
    remaining_unmet_qty: float
    realized_cost_usd: float
    non_origin_value_usd: float
    confidence_rank: int
    after_snapshot: ProductSnapshot | None = None
    rvc_gain: float | None = None
    changed: bool = False


@dataclass
class CandidateContext:
    starting_point_id: str
    sequence_no: int
    iteration_no: int
    model_code: str
    variant_id: str
    declaration_no: str


@dataclass
class BasisStats:
    custom_code_basis: str
    distinct_material_codes: set[str] = field(default_factory=set)
    total_qty: float = 0.0
    confirmed_qty: float = 0.0
    candidate_qty: float = 0.0


REVIEW_SCOPE_RANK = {
    "clean_admissible": 0,
    "date_blocked": 1,
    "ambiguity_review": 2,
    "ambiguity_date_blocked": 3,
}


def replacement_mode_choices() -> tuple[str, ...]:
    return (
        REPLACEMENT_MODE_HEURISTIC_CUSTOM_BASIS,
        REPLACEMENT_MODE_TECHNICAL_REFERENCE,
    )


def clean_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def parse_numeric(value: object) -> float:
    text = clean_text(value)
    if not text:
        return 0.0
    return float(text)


def parse_int(value: object) -> int:
    text = clean_text(value)
    if not text:
        return 0
    return int(float(text))


def to_date(value: object) -> date | None:
    text = clean_text(value)
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def looks_like_internal_material_code(text: str) -> bool:
    return bool(text and INTERNAL_MATERIAL_CODE_RE.match(text))


def normalize_reference_material_code(value: object) -> str:
    text = clean_text(value)
    if not text:
        return ""
    if re.fullmatch(r"\d+\.\d+", text):
        numeric = float(text)
        formatted = f"{numeric:.7f}"
        left, right = formatted.split(".", 1)
        return f"{left.zfill(3)}.{right}"
    return text


def load_workspace_variants(normalized_dir: Path) -> dict[str, list[object]]:
    path = normalized_dir / "bom-variants.json"
    payload = json.loads(path.read_text(encoding="utf-8"))
    variants_by_model: dict[str, list[object]] = {}
    for model_code, items in payload.items():
        variants: list[object] = []
        for item in items:
            lines = [
                SimpleNamespace(
                    material_code=clean_text(line["material_code"]),
                    qty_per_unit=parse_numeric(line["qty_per_unit"]),
                    row_no=parse_int(line.get("row_no")),
                    ordinal_key=clean_text(line.get("ordinal_key")),
                )
                for line in item["lines"]
            ]
            variants.append(
                SimpleNamespace(
                    variant_id=clean_text(item["variant_id"]),
                    model_code=clean_text(item["model_code"]),
                    bom_code=clean_text(item["bom_code"]),
                    block_index=parse_int(item.get("block_index")),
                    start_row=parse_int(item.get("start_row")),
                    end_row=parse_int(item.get("end_row")),
                    line_count=parse_int(item.get("line_count")),
                    lines=lines,
                )
            )
        variants_by_model[model_code] = variants
    return variants_by_model


def load_explicit_substitute_index(path: Path) -> dict[tuple[str, str], set[str]]:
    index: dict[tuple[str, str], set[str]] = defaultdict(set)
    if not path.exists():
        return index
    with path.open(encoding="utf-8") as handle:
        for row in csv.DictReader(handle):
            root_product_code = clean_text(row["root_product_code"])
            component_code = normalize_reference_material_code(row["component_code"])
            substitute_material_code = normalize_reference_material_code(row["substitute_material_code"])
            if not root_product_code or not component_code or not substitute_material_code:
                continue
            if component_code == substitute_material_code:
                continue
            index[(root_product_code, component_code)].add(substitute_material_code)
    return index


def resolve_custom_code_basis(row: dict[str, object]) -> str:
    declared_code = clean_text(row.get("declared_code"))
    label_code = clean_text(row.get("label_code"))
    source = clean_text(row.get("source"))
    if (
        source == "NK2_ONLY"
        and looks_like_internal_material_code(declared_code)
        and label_code
        and not looks_like_internal_material_code(label_code)
    ):
        return label_code
    return declared_code or label_code


def resolve_candidate_material_code(row: dict[str, object], custom_code_basis: str) -> str:
    confirmed_lookup_code = clean_text(row.get("confirmed_lookup_code"))
    lookup_material_code = clean_text(row.get("lookup_material_code"))
    if confirmed_lookup_code or lookup_material_code:
        return confirmed_lookup_code or lookup_material_code
    label_code = clean_text(row.get("label_code"))
    lane = label_code or custom_code_basis or "UNMAPPED"
    return f"UNMAPPED::{lane}"


def load_baseline_module() -> object:
    module_path = Path(__file__).with_name("growatt-rvc-baseline.py")
    spec = importlib.util.spec_from_file_location("growatt_rvc_baseline_runtime", module_path)
    if spec is None or spec.loader is None:
        raise SystemExit(f"Unable to load baseline module from {module_path}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def shipment_slug(shipment_id: str) -> str:
    tail = shipment_id[-4:]
    if tail.isalnum():
        return tail.lower()
    return shipment_id.lower().replace("/", "-")


def read_staff_substitutes(paths: tuple[Path, ...] = STAFF_SUBSTITUTE_PATHS) -> set[tuple[str, str]]:
    pairs: set[tuple[str, str]] = set()
    for path in paths:
        if not path.exists():
            continue
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheet = workbook[workbook.sheetnames[0]]
        for row in sheet.iter_rows(min_row=3, values_only=True):
            original = clean_text(row[0])
            candidate = clean_text(row[2])
            if not original or not candidate:
                continue
            pairs.add((original, candidate))
            pairs.add((candidate, original))
    return pairs


def load_variant_scope_index(path: Path) -> dict[tuple[str, str], dict[str, object]]:
    index: dict[tuple[str, str], dict[str, object]] = {}
    with path.open(encoding="utf-8") as handle:
        for row in csv.DictReader(handle):
            key = (
                clean_text(row["tracking_key"]),
                clean_text(row["admissible_material_code"]),
            )
            item = index.setdefault(
                key,
                {
                    "variant_ids": set(),
                    "shipment_variant_scope": clean_text(row["shipment_variant_scope"]),
                },
            )
            variant_id = clean_text(row["variant_id"])
            if variant_id:
                item["variant_ids"].add(variant_id)
            if clean_text(row["shipment_variant_scope"]):
                item["shipment_variant_scope"] = clean_text(row["shipment_variant_scope"])
    return index


def bucket_candidate_class(
    confirmed_lookup_code: str,
    mapping_status: str,
    variant_ids: tuple[str, ...],
) -> str:
    if confirmed_lookup_code:
        return "confirmed_clean"
    if mapping_status == "candidate_exact_dm_match" and variant_ids:
        return "candidate_clean"
    if mapping_status in {"candidate_not_in_dm", "no_lookup_code"}:
        return "ambiguity_review"
    return "excluded"


def build_replacement_basis(
    shared_normalized_dir: Path,
    shipment_dir: Path,
    shipment_slug_value: str | None = None,
) -> list[ReplacementBasisBucket]:
    slug_value = shipment_slug_value or shipment_dir.name
    admissibility_path = shipment_dir / "normalized" / f"{slug_value}-variant-admissibility.csv"
    scope_index = load_variant_scope_index(admissibility_path)
    buckets: list[ReplacementBasisBucket] = []

    with (shared_normalized_dir / "co-stock-tracking-updated.csv").open(encoding="utf-8") as handle:
        for row in csv.DictReader(handle):
            remaining_qty = parse_numeric(row["remaining_qty"])
            if remaining_qty <= 0:
                continue
            custom_code_basis = resolve_custom_code_basis(row)
            candidate_material_code = resolve_candidate_material_code(row, custom_code_basis)
            erp_material_code = clean_text(row["confirmed_lookup_code"]) or clean_text(row["lookup_material_code"]) or candidate_material_code
            if not custom_code_basis or not candidate_material_code:
                continue
            tracking_key = clean_text(row["tracking_key"])
            scope = scope_index.get((tracking_key, erp_material_code), {})
            variant_ids = tuple(sorted(scope.get("variant_ids", set())))
            candidate_class = bucket_candidate_class(
                clean_text(row["confirmed_lookup_code"]),
                clean_text(row["mapping_status"]),
                variant_ids,
            )
            source_row_no = clean_text(row["nk2_source_row_no"]) or clean_text(row["bcct_source_row_no"]) or "na"
            buckets.append(
                ReplacementBasisBucket(
                    custom_code_basis=custom_code_basis,
                    erp_material_code=erp_material_code,
                    bucket_id=f"{tracking_key}-{erp_material_code}-{source_row_no}",
                    tracking_key=tracking_key,
                    declaration_no=clean_text(row["declaration_no"]),
                    declaration_item_no=parse_int(row["declaration_item_no"]),
                    import_date=to_date(row["declaration_date"]),
                    source=clean_text(row["source"]),
                    source_row_no=source_row_no,
                    hs_code=clean_text(row["hs_code"]),
                    name=clean_text(row["name"]),
                    origin=clean_text(row["origin"]),
                    unit_price_usd=parse_numeric(row["unit_price_usd"]),
                    exchange_rate=parse_numeric(row["exchange_rate"]),
                    remaining_qty=remaining_qty,
                    candidate_class=candidate_class,
                    mapping_status=clean_text(row["mapping_status"]),
                    mapping_confidence=clean_text(row["lookup_confidence"]),
                    evidence_source=clean_text(row["mapping_status"]),
                    shipment_variant_scope=clean_text(scope.get("shipment_variant_scope")),
                    variant_scope=variant_ids,
                    variant_hit_count=len(variant_ids),
                    candidate_material_code=candidate_material_code,
                )
            )
    return buckets


def clone_stock_indexes(
    basis_buckets: list[ReplacementBasisBucket],
) -> tuple[dict[str, list[ReplacementBasisBucket]], dict[str, list[ReplacementBasisBucket]], dict[str, ReplacementBasisBucket]]:
    by_material: dict[str, list[ReplacementBasisBucket]] = defaultdict(list)
    by_basis: dict[str, list[ReplacementBasisBucket]] = defaultdict(list)
    by_id: dict[str, ReplacementBasisBucket] = {}
    for bucket in basis_buckets:
        if bucket.candidate_class == "excluded":
            continue
        cloned = copy.copy(bucket)
        by_material[cloned.erp_material_code].append(cloned)
        by_basis[cloned.custom_code_basis].append(cloned)
        by_id[cloned.bucket_id] = cloned
    return by_material, by_basis, by_id


def build_material_basis_stats(
    basis_buckets: list[ReplacementBasisBucket],
) -> dict[str, dict[str, BasisStats]]:
    stats_by_material: dict[str, dict[str, BasisStats]] = defaultdict(dict)
    for bucket in basis_buckets:
        if bucket.candidate_class == "excluded":
            continue
        by_basis = stats_by_material[bucket.erp_material_code]
        stats = by_basis.setdefault(bucket.custom_code_basis, BasisStats(custom_code_basis=bucket.custom_code_basis))
        stats.distinct_material_codes.add(bucket.erp_material_code)
        stats.total_qty += bucket.remaining_qty
        if bucket.candidate_class == "confirmed_clean":
            stats.confirmed_qty += bucket.remaining_qty
        if bucket.candidate_class == "candidate_clean":
            stats.candidate_qty += bucket.remaining_qty
    return stats_by_material


def choose_custom_code_basis(
    original_material_code: str,
    stats_by_material: dict[str, dict[str, BasisStats]],
    stock_by_basis: dict[str, list[ReplacementBasisBucket]],
) -> str:
    by_basis = stats_by_material.get(original_material_code, {})
    if not by_basis:
        return original_material_code

    def sort_key(stats: BasisStats) -> tuple[int, int, float, int, float, str]:
        basis_buckets = stock_by_basis.get(stats.custom_code_basis, [])
        distinct_codes = {item.erp_material_code for item in basis_buckets if item.remaining_qty > EPS}
        return (
            1 if len(distinct_codes) > 1 else 0,
            len(distinct_codes),
            sum(item.remaining_qty for item in basis_buckets if item.remaining_qty > EPS),
            1 if stats.custom_code_basis != original_material_code else 0,
            stats.confirmed_qty,
            stats.custom_code_basis,
        )

    ranked = sorted(by_basis.values(), key=sort_key, reverse=True)
    return ranked[0].custom_code_basis


def confidence_rank(value: str) -> int:
    lookup = {
        "medium": 3,
        "candidate": 2,
        "low": 1,
        "none": 0,
    }
    return lookup.get(value, 0)


def candidate_class_rank(value: str) -> int:
    lookup = {
        "confirmed_clean": 2,
        "candidate_clean": 1,
        "ambiguity_review": 0,
        "excluded": 0,
    }
    return lookup.get(value, 0)


def allocation_sort_key(baseline: object, bucket: ReplacementBasisBucket) -> tuple[float, float, int, int, object, str]:
    non_origin_unit_cost = 0.0 if baseline.is_vietnam_origin(bucket.origin) else bucket.unit_price_usd
    return (
        non_origin_unit_cost,
        bucket.unit_price_usd,
        -candidate_class_rank(bucket.candidate_class),
        -confidence_rank(bucket.mapping_confidence),
        bucket.import_date or date.max,
        bucket.bucket_id,
    )


def bucket_is_date_eligible(
    baseline: object,
    bucket: ReplacementBasisBucket,
    export_line: object,
    import_lead_days: int,
    max_import_age_days: int,
) -> bool:
    return baseline.is_bucket_eligible_for_export(
        bucket,
        export_line,
        import_lead_days=import_lead_days,
        max_import_age_days=max_import_age_days,
    )


def bucket_is_admissible_for_variant(bucket: ReplacementBasisBucket, variant_id: str) -> bool:
    if bucket.candidate_class == "candidate_clean":
        return variant_id in bucket.variant_scope
    if not bucket.variant_scope:
        return True
    return variant_id in bucket.variant_scope


def bucket_review_scope(
    baseline: object,
    bucket: ReplacementBasisBucket,
    export_line: object,
    variant_id: str,
    import_lead_days: int,
    max_import_age_days: int,
) -> tuple[bool, str, tuple[str, ...]]:
    if bucket.remaining_qty <= EPS or bucket.candidate_class == "excluded":
        return False, "", ()

    risk_flags: list[str] = []
    if bucket.candidate_class in {"confirmed_clean", "candidate_clean"}:
        if not bucket_is_admissible_for_variant(bucket, variant_id):
            return False, "", ()
    else:
        risk_flags.append("ambiguity")
        if bucket.mapping_status == "no_lookup_code":
            risk_flags.append("no_lookup_code")

    if not bucket_is_date_eligible(
        baseline,
        bucket,
        export_line,
        import_lead_days,
        max_import_age_days,
    ):
        risk_flags.append("date_blocked")

    if not risk_flags:
        return True, "clean_admissible", ()
    if risk_flags == ["date_blocked"]:
        return True, "date_blocked", ("date_blocked",)
    if risk_flags == ["ambiguity"]:
        return True, "ambiguity_review", ("ambiguity",)
    return True, "ambiguity_date_blocked", ("ambiguity", "date_blocked")


def review_allocation_sort_key(
    baseline: object,
    bucket: ReplacementBasisBucket,
    review_scope: str,
) -> tuple[int, float, float, int, int, object, str]:
    base_key = allocation_sort_key(baseline, bucket)
    return (REVIEW_SCOPE_RANK[review_scope], *base_key)


def bucket_is_eligible(
    baseline: object,
    bucket: ReplacementBasisBucket,
    export_line: object,
    variant_id: str,
    import_lead_days: int,
    max_import_age_days: int,
) -> bool:
    if bucket.remaining_qty <= EPS:
        return False
    if bucket.candidate_class == "excluded":
        return False
    if bucket.candidate_class in {"confirmed_clean", "candidate_clean"} and not bucket_is_admissible_for_variant(
        bucket,
        variant_id,
    ):
        return False
    return bucket_is_date_eligible(
        baseline,
        bucket,
        export_line,
        import_lead_days,
        max_import_age_days,
    )


def material_state_clone(materials: list[MaterialState]) -> list[MaterialState]:
    return copy.deepcopy(materials)


def summarize_product(baseline: object, export_line: object, materials: list[MaterialState]) -> ProductSnapshot:
    non_origin_value = 0.0
    unmet_qty_total = 0.0
    for material in materials:
        unmet_qty_total += max(material.unmet_qty, 0.0)
        for allocation in material.exact_allocations + material.replacement_allocations:
            if baseline.is_vietnam_origin(allocation.origin):
                continue
            non_origin_value += allocation.allocated_qty * allocation.unit_price_usd

    fob_value_usd = export_line.quantity * export_line.unit_price_usd
    rvc_percent = None
    if fob_value_usd > 0:
        rvc_percent = ((fob_value_usd - non_origin_value) / fob_value_usd) * 100.0
    margin = None if rvc_percent is None else rvc_percent - baseline.TARGET_RVC
    stock_sufficient = unmet_qty_total <= EPS
    passes_rvc = stock_sufficient and rvc_percent is not None and rvc_percent >= baseline.TARGET_RVC
    return ProductSnapshot(
        rvc_percent=rvc_percent,
        margin_to_threshold=margin,
        stock_sufficient=stock_sufficient,
        passes_rvc=passes_rvc,
        unmet_qty_total=unmet_qty_total,
        non_origin_value_usd=non_origin_value,
        materials=materials,
    )


def current_material_non_origin_value(baseline: object, material: MaterialState) -> float:
    return sum(
        item.allocated_qty * item.unit_price_usd
        for item in material.active_allocations
        if not baseline.is_vietnam_origin(item.origin)
    )


def build_virtual_basis_buckets(
    target_material: MaterialState,
    stock_by_basis: dict[str, list[ReplacementBasisBucket]],
) -> list[ReplacementBasisBucket]:
    released_qty_by_bucket = defaultdict(float)
    for allocation in target_material.active_allocations:
        released_qty_by_bucket[allocation.bucket_id] += allocation.allocated_qty

    virtual_buckets: list[ReplacementBasisBucket] = []
    custom_code_basis = target_material.custom_code_basis or target_material.original_material_code
    for bucket in stock_by_basis.get(custom_code_basis, []):
        clone = copy.copy(bucket)
        clone.remaining_qty += released_qty_by_bucket.get(bucket.bucket_id, 0.0)
        virtual_buckets.append(clone)
    return virtual_buckets


def build_virtual_material_buckets(
    target_material: MaterialState,
    stock_by_material: dict[str, list[ReplacementBasisBucket]],
    allowed_material_codes: set[str],
) -> list[ReplacementBasisBucket]:
    released_qty_by_bucket = defaultdict(float)
    for allocation in target_material.active_allocations:
        released_qty_by_bucket[allocation.bucket_id] += allocation.allocated_qty

    virtual_buckets: list[ReplacementBasisBucket] = []
    for material_code in sorted(allowed_material_codes):
        for bucket in stock_by_material.get(material_code, []):
            clone = copy.copy(bucket)
            clone.remaining_qty += released_qty_by_bucket.get(bucket.bucket_id, 0.0)
            virtual_buckets.append(clone)
    return virtual_buckets


def build_candidate_pool(
    target_material: MaterialState,
    export_line: object,
    stock_by_material: dict[str, list[ReplacementBasisBucket]],
    stock_by_basis: dict[str, list[ReplacementBasisBucket]],
    replacement_mode: str,
    explicit_substitute_index: dict[tuple[str, str], set[str]] | None,
) -> tuple[str, list[ReplacementBasisBucket], set[str]]:
    custom_code_basis = target_material.custom_code_basis or target_material.original_material_code
    if replacement_mode == REPLACEMENT_MODE_TECHNICAL_REFERENCE:
        explicit_candidates = set(
            (explicit_substitute_index or {}).get(
                (export_line.model_code, target_material.original_material_code),
                set(),
            )
        )
        allowed_material_codes = {target_material.original_material_code, *explicit_candidates}
        return (
            custom_code_basis,
            build_virtual_material_buckets(
                target_material,
                stock_by_material,
                allowed_material_codes,
            ),
            explicit_candidates,
        )
    return (
        custom_code_basis,
        build_virtual_basis_buckets(target_material, stock_by_basis),
        set(),
    )


def allocations_changed(material: MaterialState, allocations: list[BucketAllocation]) -> bool:
    current = [
        (
            item.bucket_id,
            item.material_code,
            round(item.allocated_qty, 9),
        )
        for item in material.active_allocations
    ]
    proposed = [
        (
            item.bucket_id,
            item.material_code,
            round(item.allocated_qty, 9),
        )
        for item in allocations
    ]
    return current != proposed


def allocate_from_bucket_list(
    buckets: list[ReplacementBasisBucket],
    required_qty: float,
    custom_code_basis: str,
    staff_substitute: bool,
    *,
    mutate: bool,
) -> list[BucketAllocation]:
    allocations: list[BucketAllocation] = []
    remaining_need = required_qty
    for bucket in buckets:
        if remaining_need <= EPS:
            break
        available_qty = bucket.remaining_qty
        if available_qty <= EPS:
            continue
        take_qty = min(available_qty, remaining_need)
        if take_qty <= EPS:
            continue
        if mutate:
            bucket.remaining_qty -= take_qty
        remaining_need -= take_qty
        allocations.append(
            BucketAllocation(
                bucket_id=bucket.bucket_id,
                tracking_key=bucket.tracking_key,
                material_code=bucket.erp_material_code,
                custom_code_basis=custom_code_basis,
                declaration_no=bucket.declaration_no,
                declaration_item_no=bucket.declaration_item_no,
                import_date=bucket.import_date,
                allocated_qty=take_qty,
                unit_price_usd=bucket.unit_price_usd,
                exchange_rate=bucket.exchange_rate,
                origin=bucket.origin,
                source=bucket.source,
                source_row_no=bucket.source_row_no,
                candidate_class=bucket.candidate_class,
                mapping_status=bucket.mapping_status,
                mapping_confidence=bucket.mapping_confidence,
                evidence_source=bucket.evidence_source,
                staff_known_substitute=staff_substitute,
            )
        )
    return allocations


def allocate_exact_bom(
    baseline: object,
    export_line: object,
    variant: object,
    stock_by_material: dict[str, list[ReplacementBasisBucket]],
    material_basis_stats: dict[str, dict[str, BasisStats]],
    stock_by_basis: dict[str, list[ReplacementBasisBucket]],
    import_lead_days: int,
    max_import_age_days: int,
) -> list[MaterialState]:
    materials: list[MaterialState] = []
    for bom_line in variant.lines:
        need_qty = bom_line.qty_per_unit * export_line.quantity
        custom_code_basis = choose_custom_code_basis(
            bom_line.material_code,
            material_basis_stats,
            stock_by_basis,
        )
        eligible_buckets = [
            bucket
            for bucket in stock_by_material.get(bom_line.material_code, [])
            if bucket_is_eligible(
                baseline,
                bucket,
                export_line,
                variant.variant_id,
                import_lead_days,
                max_import_age_days,
            )
        ]
        allocations = allocate_from_bucket_list(
            eligible_buckets,
            required_qty=need_qty,
            custom_code_basis=custom_code_basis,
            staff_substitute=False,
            mutate=True,
        )
        allocated_qty = sum(item.allocated_qty for item in allocations)
        materials.append(
            MaterialState(
                original_material_code=bom_line.material_code,
                need_qty=need_qty,
                exact_allocations=allocations,
                exact_allocated_qty=allocated_qty,
                unmet_qty=max(0.0, need_qty - allocated_qty),
                custom_code_basis=custom_code_basis,
            )
        )
    return materials


def preview_snapshot_after_allocations(
    baseline: object,
    export_line: object,
    current_snapshot: ProductSnapshot,
    target_material: MaterialState,
    allocations: list[BucketAllocation],
) -> ProductSnapshot:
    new_allocated_qty = sum(item.allocated_qty for item in allocations)
    new_unmet_qty = max(0.0, target_material.need_qty - new_allocated_qty)
    new_non_origin_value = sum(
        item.allocated_qty * item.unit_price_usd
        for item in allocations
        if not baseline.is_vietnam_origin(item.origin)
    )
    unmet_qty_total = max(0.0, current_snapshot.unmet_qty_total - target_material.unmet_qty + new_unmet_qty)
    non_origin_value = (
        current_snapshot.non_origin_value_usd
        - current_material_non_origin_value(baseline, target_material)
        + new_non_origin_value
    )
    fob_value_usd = export_line.quantity * export_line.unit_price_usd
    rvc_percent = None
    if fob_value_usd > 0:
        rvc_percent = ((fob_value_usd - non_origin_value) / fob_value_usd) * 100.0
    margin = None if rvc_percent is None else rvc_percent - baseline.TARGET_RVC
    stock_sufficient = unmet_qty_total <= EPS
    passes_rvc = stock_sufficient and rvc_percent is not None and rvc_percent >= baseline.TARGET_RVC
    return ProductSnapshot(
        rvc_percent=rvc_percent,
        margin_to_threshold=margin,
        stock_sufficient=stock_sufficient,
        passes_rvc=passes_rvc,
        unmet_qty_total=unmet_qty_total,
        non_origin_value_usd=non_origin_value,
        materials=current_snapshot.materials,
    )


def build_material_plan(
    baseline: object,
    export_line: object,
    variant: object,
    current_materials: list[MaterialState],
    current_snapshot: ProductSnapshot,
    target_material: MaterialState,
    stock_by_basis: dict[str, list[ReplacementBasisBucket]],
    staff_substitutes: set[tuple[str, str]],
    import_lead_days: int,
    max_import_age_days: int,
    stock_by_material: dict[str, list[ReplacementBasisBucket]] | None = None,
    replacement_mode: str = REPLACEMENT_MODE_HEURISTIC_CUSTOM_BASIS,
    explicit_substitute_index: dict[tuple[str, str], set[str]] | None = None,
) -> MaterialPlan | None:
    custom_code_basis, virtual_buckets, _ = build_candidate_pool(
        target_material,
        export_line,
        stock_by_material or {},
        stock_by_basis,
        replacement_mode,
        explicit_substitute_index,
    )
    eligible_buckets = [
        bucket
        for bucket in virtual_buckets
        if bucket_is_eligible(
            baseline,
            bucket,
            export_line,
            variant.variant_id,
            import_lead_days,
            max_import_age_days,
        )
    ]
    if not eligible_buckets:
        return None

    ranked_buckets = sorted(eligible_buckets, key=lambda item: allocation_sort_key(baseline, item))
    allocations: list[BucketAllocation] = []
    remaining_need = target_material.need_qty
    for bucket in ranked_buckets:
        if remaining_need <= EPS:
            break
        preview = allocate_from_bucket_list(
            [bucket],
            required_qty=remaining_need,
            custom_code_basis=custom_code_basis,
            staff_substitute=(target_material.original_material_code, bucket.erp_material_code) in staff_substitutes,
            mutate=False,
        )
        if not preview:
            continue
        allocations.extend(preview)
        remaining_need -= sum(item.allocated_qty for item in preview)
    allocated_qty = sum(item.allocated_qty for item in allocations)
    if allocated_qty <= EPS or remaining_need > EPS:
        return None

    realized_cost = sum(item.allocated_qty * item.unit_price_usd for item in allocations)
    non_origin_value = sum(
        item.allocated_qty * item.unit_price_usd
        for item in allocations
        if not baseline.is_vietnam_origin(item.origin)
    )
    confidence = max((confidence_rank(item.mapping_confidence) for item in allocations), default=0)
    after_snapshot = preview_snapshot_after_allocations(
        baseline,
        export_line,
        current_snapshot,
        target_material,
        allocations,
    )
    rvc_gain = None
    if current_snapshot.rvc_percent is not None and after_snapshot.rvc_percent is not None:
        rvc_gain = after_snapshot.rvc_percent - current_snapshot.rvc_percent
    changed = allocations_changed(target_material, allocations) or abs(
        max(0.0, target_material.need_qty - allocated_qty) - target_material.unmet_qty
    ) > EPS

    return MaterialPlan(
        original_material_code=target_material.original_material_code,
        custom_code_basis=custom_code_basis,
        allocations=allocations,
        covered_qty=allocated_qty,
        remaining_unmet_qty=max(0.0, target_material.need_qty - allocated_qty),
        realized_cost_usd=realized_cost,
        non_origin_value_usd=non_origin_value,
        confidence_rank=confidence,
        after_snapshot=after_snapshot,
        rvc_gain=rvc_gain,
        changed=changed,
    )


def apply_material_plan(
    material: MaterialState,
    plan: MaterialPlan,
    bucket_by_id: dict[str, ReplacementBasisBucket],
) -> None:
    for allocation in material.active_allocations:
        bucket_by_id[allocation.bucket_id].remaining_qty += allocation.allocated_qty
    material.exact_allocations = []
    material.replacement_allocations = []
    material.exact_allocated_qty = 0.0
    for allocation in plan.allocations:
        bucket = bucket_by_id[allocation.bucket_id]
        bucket.remaining_qty -= allocation.allocated_qty
        material.replacement_allocations.append(copy.deepcopy(allocation))
    material.unmet_qty = plan.remaining_unmet_qty


def group_candidate_allocations(
    allocations: list[BucketAllocation],
) -> dict[str, list[BucketAllocation]]:
    grouped: dict[str, list[BucketAllocation]] = defaultdict(list)
    for item in allocations:
        grouped[item.material_code].append(item)
    return grouped


def candidate_sort_key(row: dict[str, object]) -> tuple[int, float, float, int, str]:
    return (
        1 if row["would_reach_target"] else 0,
        1 if row["commit_scope_eligible"] else 0,
        -int(row["risk_flag_count"]),
        float(row["estimated_rvc_after"] if row["estimated_rvc_after"] is not None else -9999.0),
        -float(row["realized_cost_usd"]),
        int(row["confidence_rank"]),
        str(row["candidate_material_code"]),
    )


def build_candidate_rows(
    baseline: object,
    context: CandidateContext,
    export_line: object,
    variant: object,
    current_materials: list[MaterialState],
    current_snapshot: ProductSnapshot,
    target_material: MaterialState,
    stock_by_basis: dict[str, list[ReplacementBasisBucket]],
    staff_substitutes: set[tuple[str, str]],
    import_lead_days: int,
    max_import_age_days: int,
    stock_by_material: dict[str, list[ReplacementBasisBucket]] | None = None,
    replacement_mode: str = REPLACEMENT_MODE_HEURISTIC_CUSTOM_BASIS,
    explicit_substitute_index: dict[tuple[str, str], set[str]] | None = None,
) -> list[dict[str, object]]:
    custom_code_basis, virtual_buckets, explicit_candidates = build_candidate_pool(
        target_material,
        export_line,
        stock_by_material or {},
        stock_by_basis,
        replacement_mode,
        explicit_substitute_index,
    )
    by_candidate: dict[str, list[tuple[ReplacementBasisBucket, str, tuple[str, ...]]]] = defaultdict(list)
    for bucket in virtual_buckets:
        candidate_material_code = bucket.candidate_material_code or bucket.erp_material_code
        if candidate_material_code == target_material.original_material_code:
            continue
        if replacement_mode == REPLACEMENT_MODE_TECHNICAL_REFERENCE and candidate_material_code not in explicit_candidates:
            continue
        include, review_scope, risk_flags = bucket_review_scope(
            baseline,
            bucket,
            export_line,
            variant.variant_id,
            import_lead_days,
            max_import_age_days,
        )
        if not include:
            continue
        by_candidate[candidate_material_code].append((bucket, review_scope, risk_flags))

    rows: list[dict[str, object]] = []
    for candidate_material_code, bucket_entries in by_candidate.items():
        buckets = [item[0] for item in bucket_entries]
        review_scope_by_bucket = {item[0].bucket_id: item[1] for item in bucket_entries}
        risk_flags_by_bucket = {item[0].bucket_id: item[2] for item in bucket_entries}
        preview_allocations = allocate_from_bucket_list(
            [
                item[0]
                for item in sorted(
                    bucket_entries,
                    key=lambda item: review_allocation_sort_key(baseline, item[0], item[1]),
                )
            ],
            required_qty=target_material.need_qty,
            custom_code_basis=custom_code_basis,
            staff_substitute=(target_material.original_material_code, candidate_material_code) in staff_substitutes,
            mutate=False,
        )
        if not preview_allocations:
            continue
        admissible_qty = sum(item.remaining_qty for item in buckets)
        scope_qtys: dict[str, float] = defaultdict(float)
        risk_flags: set[str] = set()
        for bucket, review_scope, bucket_risk_flags in bucket_entries:
            scope_qtys[review_scope] += bucket.remaining_qty
            risk_flags.update(bucket_risk_flags)
        allocated_preview_qty = sum(item.allocated_qty for item in preview_allocations)
        fully_covered = target_material.need_qty - allocated_preview_qty <= EPS
        after_snapshot = preview_snapshot_after_allocations(
            baseline,
            export_line,
            current_snapshot,
            target_material,
            preview_allocations,
        )
        estimated_gain = None
        if current_snapshot.rvc_percent is not None and after_snapshot.rvc_percent is not None:
            estimated_gain = after_snapshot.rvc_percent - current_snapshot.rvc_percent
        realized_cost = sum(item.allocated_qty * item.unit_price_usd for item in preview_allocations)
        used_review_scopes = sorted({review_scope_by_bucket[item.bucket_id] for item in preview_allocations})
        used_risk_flags = sorted(
            {
                flag
                for item in preview_allocations
                for flag in risk_flags_by_bucket[item.bucket_id]
            }
        )
        rows.append(
            {
                "starting_point_id": context.starting_point_id,
                "sequence_no": context.sequence_no,
                "iteration_no": context.iteration_no,
                "model_code": context.model_code,
                "variant_id": context.variant_id,
                "declaration_no": context.declaration_no,
                "original_material_code": target_material.original_material_code,
                "custom_code_basis": custom_code_basis,
                "candidate_material_code": candidate_material_code,
                "admissible_qty": admissible_qty,
                "allocated_preview_qty": allocated_preview_qty,
                "bucket_cost_basis": json.dumps(
                    [
                        {
                            "bucket_id": item.bucket_id,
                            "allocated_qty": item.allocated_qty,
                            "unit_price_usd": item.unit_price_usd,
                            "origin": item.origin,
                            "review_scope": review_scope_by_bucket[item.bucket_id],
                            "risk_flags": list(risk_flags_by_bucket[item.bucket_id]),
                        }
                        for item in preview_allocations
                    ],
                    ensure_ascii=False,
                    default=str,
                ),
                "mapping_confidence": max(
                    (item.mapping_confidence for item in preview_allocations),
                    key=confidence_rank,
                ),
                "confidence_rank": max(
                    (confidence_rank(item.mapping_confidence) for item in preview_allocations),
                    default=0,
                ),
                "mapping_evidence_source": ";".join(
                    sorted({item.evidence_source for item in preview_allocations if item.evidence_source})
                ),
                "staff_known_substitute": any(item.staff_known_substitute for item in preview_allocations),
                "estimated_rvc_after": after_snapshot.rvc_percent,
                "estimated_rvc_gain": estimated_gain,
                "would_reach_target": after_snapshot.passes_rvc,
                "realized_cost_usd": realized_cost,
                "variant_hit_scope": ";".join(
                    sorted(
                        {
                            scope_id
                            for bucket in buckets
                            for scope_id in bucket.variant_scope
                        }
                    )
                ),
                "candidate_class": ";".join(sorted({item.candidate_class for item in preview_allocations})),
                "review_scope": "mixed" if len(used_review_scopes) > 1 else used_review_scopes[0],
                "risk_flags": ";".join(used_risk_flags),
                "risk_flag_count": len(used_risk_flags),
                "commit_scope_eligible": fully_covered and "date_blocked" not in used_risk_flags,
                "estimate_scope": "same_custom_code_commit" if "date_blocked" not in used_risk_flags else "review_only_date_blocked",
                "clean_qty_available": scope_qtys.get("clean_admissible", 0.0),
                "date_blocked_qty_available": scope_qtys.get("date_blocked", 0.0),
                "ambiguity_qty_available": scope_qtys.get("ambiguity_review", 0.0),
                "ambiguity_date_blocked_qty_available": scope_qtys.get("ambiguity_date_blocked", 0.0),
                "remaining_unmet_after_candidate": max(0.0, target_material.need_qty - allocated_preview_qty),
            }
        )

    rows.sort(key=candidate_sort_key, reverse=True)
    for idx, row in enumerate(rows, start=1):
        row["candidate_rank"] = idx
    return rows


def snapshot_objective_key(snapshot: ProductSnapshot | None) -> tuple[int, float, float, float]:
    if snapshot is None:
        return (-1, -9999.0, -9999.0, -9999.0)
    return (
        1 if snapshot.passes_rvc else 0,
        1 if snapshot.stock_sufficient else 0,
        -float(snapshot.unmet_qty_total),
        -float(snapshot.non_origin_value_usd),
    )


def plan_sort_key(plan: MaterialPlan) -> tuple[int, float, float, float, int, int]:
    after = plan.after_snapshot
    if after is None:
        return (-1, -9999.0, -9999.0, -9999.0, -9999, 0)
    return (
        *snapshot_objective_key(after),
        1 if plan.changed else 0,
        int(plan.confidence_rank),
    )


def plan_improves(plan: MaterialPlan, current_snapshot: ProductSnapshot) -> bool:
    if not plan.changed:
        return False
    after = plan.after_snapshot
    if after is None:
        return False
    if after.passes_rvc != current_snapshot.passes_rvc:
        return after.passes_rvc
    if after.stock_sufficient != current_snapshot.stock_sufficient:
        return after.stock_sufficient
    if after.unmet_qty_total < current_snapshot.unmet_qty_total - EPS:
        return True
    if after.unmet_qty_total > current_snapshot.unmet_qty_total + EPS:
        return False
    if after.non_origin_value_usd < current_snapshot.non_origin_value_usd - COST_EPS:
        return True
    if after.non_origin_value_usd > current_snapshot.non_origin_value_usd + COST_EPS:
        return False
    return False


def snapshot_to_payload(snapshot: ProductSnapshot) -> dict[str, object]:
    return {
        "rvc_percent": snapshot.rvc_percent,
        "margin_to_threshold": snapshot.margin_to_threshold,
        "stock_sufficient": snapshot.stock_sufficient,
        "passes_rvc": snapshot.passes_rvc,
        "unmet_qty_total": snapshot.unmet_qty_total,
        "non_origin_value_usd": snapshot.non_origin_value_usd,
        "materials": [
            {
                "original_material_code": material.original_material_code,
                "custom_code_basis": material.custom_code_basis,
                "need_qty": material.need_qty,
                "exact_allocated_qty": material.exact_allocated_qty,
                "replacement_allocated_qty": material.replacement_allocated_qty,
                "unmet_qty": material.unmet_qty,
                "exact_allocations": [asdict(item) for item in material.exact_allocations],
                "replacement_allocations": [asdict(item) for item in material.replacement_allocations],
            }
            for material in snapshot.materials
        ],
    }


def stock_state_summary(stock_by_material: dict[str, list[ReplacementBasisBucket]]) -> dict[str, object]:
    by_material = []
    for material_code, buckets in sorted(stock_by_material.items()):
        remaining_qty = sum(item.remaining_qty for item in buckets if item.remaining_qty > EPS)
        if remaining_qty <= EPS:
            continue
        by_material.append(
            {
                "material_code": material_code,
                "remaining_qty": remaining_qty,
                "bucket_count": sum(1 for item in buckets if item.remaining_qty > EPS),
            }
        )
    return {"remaining_qty_by_material_code": by_material}


def stock_bucket_rows(
    stock_by_material: dict[str, list[ReplacementBasisBucket]],
    *,
    include_zero: bool = True,
) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for material_code, buckets in sorted(stock_by_material.items()):
        for bucket in sorted(buckets, key=lambda item: (item.bucket_id, item.tracking_key)):
            if not include_zero and bucket.remaining_qty <= EPS:
                continue
            rows.append(
                {
                    "material_code": material_code,
                    "bucket_material_code": bucket.erp_material_code,
                    "candidate_material_code": bucket.candidate_material_code,
                    "custom_code_basis": bucket.custom_code_basis,
                    "bucket_id": bucket.bucket_id,
                    "tracking_key": bucket.tracking_key,
                    "declaration_no": bucket.declaration_no,
                    "declaration_item_no": bucket.declaration_item_no,
                    "import_date": bucket.import_date,
                    "source": bucket.source,
                    "source_row_no": bucket.source_row_no,
                    "hs_code": bucket.hs_code,
                    "name": bucket.name,
                    "origin": bucket.origin,
                    "unit_price_usd": bucket.unit_price_usd,
                    "exchange_rate": bucket.exchange_rate,
                    "remaining_qty": bucket.remaining_qty,
                    "candidate_class": bucket.candidate_class,
                    "mapping_status": bucket.mapping_status,
                    "mapping_confidence": bucket.mapping_confidence,
                    "evidence_source": bucket.evidence_source,
                    "shipment_variant_scope": bucket.shipment_variant_scope,
                    "variant_scope": ";".join(bucket.variant_scope),
                    "variant_hit_count": bucket.variant_hit_count,
                }
            )
    return rows


def stock_material_summary_rows(stock_by_material: dict[str, list[ReplacementBasisBucket]]) -> list[dict[str, object]]:
    return stock_state_summary(stock_by_material)["remaining_qty_by_material_code"]


def product_consumption_rows(
    starting_point_id: str,
    sequence_no: int,
    export_line: object,
    variant: object,
    materials: list[MaterialState],
) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for material in materials:
        allocation_groups = (
            ("exact", material.exact_allocations),
            ("replacement", material.replacement_allocations),
        )
        for allocation_type, allocations in allocation_groups:
            for allocation in allocations:
                rows.append(
                    {
                        "starting_point_id": starting_point_id,
                        "sequence_no": sequence_no,
                        "model_code": export_line.model_code,
                        "variant_id": variant.variant_id,
                        "bom_code": variant.bom_code,
                        "declaration_no": export_line.declaration_no,
                        "original_material_code": material.original_material_code,
                        "allocation_type": allocation_type,
                        "allocated_material_code": allocation.material_code,
                        "custom_code_basis": allocation.custom_code_basis,
                        "bucket_id": allocation.bucket_id,
                        "tracking_key": allocation.tracking_key,
                        "bucket_declaration_no": allocation.declaration_no,
                        "bucket_declaration_item_no": allocation.declaration_item_no,
                        "import_date": allocation.import_date,
                        "allocated_qty": allocation.allocated_qty,
                        "unit_price_usd": allocation.unit_price_usd,
                        "allocated_value_usd": allocation.allocated_qty * allocation.unit_price_usd,
                        "origin": allocation.origin,
                        "source": allocation.source,
                        "source_row_no": allocation.source_row_no,
                        "candidate_class": allocation.candidate_class,
                        "mapping_status": allocation.mapping_status,
                        "mapping_confidence": allocation.mapping_confidence,
                        "evidence_source": allocation.evidence_source,
                        "staff_known_substitute": allocation.staff_known_substitute,
                    }
                )
    rows.sort(
        key=lambda item: (
            item["sequence_no"],
            item["original_material_code"],
            item["allocation_type"],
            item["allocated_material_code"],
            item["bucket_id"],
        )
    )
    return rows


def stock_after_product_rows(
    starting_point_id: str,
    sequence_no: int,
    export_line: object,
    summary: dict[str, object],
) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for item in summary["remaining_qty_by_material_code"]:
        rows.append(
            {
                "starting_point_id": starting_point_id,
                "sequence_no": sequence_no,
                "model_code": export_line.model_code,
                "declaration_no": export_line.declaration_no,
                "material_code": item["material_code"],
                "remaining_qty": item["remaining_qty"],
                "bucket_count": item["bucket_count"],
            }
        )
    return rows


def candidate_row_key(row: dict[str, object]) -> tuple[object, ...]:
    return (
        row["starting_point_id"],
        row["sequence_no"],
        row["model_code"],
        row["variant_id"],
        row["declaration_no"],
        row["original_material_code"],
        row["candidate_material_code"],
    )


def product_state_signature(materials: list[MaterialState]) -> tuple[object, ...]:
    return tuple(
        (
            material.original_material_code,
            round(material.unmet_qty, 9),
            tuple(
                (
                    allocation.bucket_id,
                    allocation.material_code,
                    round(allocation.allocated_qty, 9),
                )
                for allocation in material.active_allocations
            ),
        )
        for material in materials
    )


def run_starting_point(
    baseline: object,
    shipment_id: str,
    starting_point_id: str,
    export_lines: list[object],
    variant_by_id: dict[str, object],
    initial_stock_by_material: dict[str, list[ReplacementBasisBucket]],
    material_basis_stats: dict[str, dict[str, BasisStats]],
    staff_substitutes: set[tuple[str, str]],
    import_lead_days: int,
    max_import_age_days: int,
    output_dir: Path,
    replacement_mode: str = REPLACEMENT_MODE_HEURISTIC_CUSTOM_BASIS,
    explicit_substitute_index: dict[tuple[str, str], set[str]] | None = None,
    collect_candidate_rows: bool = True,
) -> dict[str, object]:
    stock_by_material = {
        code: [copy.copy(bucket) for bucket in buckets]
        for code, buckets in initial_stock_by_material.items()
    }
    bucket_by_id: dict[str, ReplacementBasisBucket] = {}
    stock_by_basis: dict[str, list[ReplacementBasisBucket]] = defaultdict(list)
    for code, buckets in stock_by_material.items():
        for bucket in buckets:
            bucket_by_id[bucket.bucket_id] = bucket
            stock_by_basis[bucket.custom_code_basis].append(bucket)

    candidate_rows_by_key: dict[tuple[object, ...], dict[str, object]] = {}
    product_status_rows: list[dict[str, object]] = []
    initial_bucket_rows = stock_bucket_rows(stock_by_material)
    consumption_rows: list[dict[str, object]] = []
    after_product_stock_rows: list[dict[str, object]] = []
    snapshot_dir = output_dir / "replacement-bom-snapshots" / starting_point_id
    snapshot_dir.mkdir(parents=True, exist_ok=True)
    final_products: list[dict[str, object]] = []
    changed_products: list[str] = []
    changed_materials: set[str] = set()

    for sequence_no, export_line in enumerate(export_lines, start=1):
        variant = variant_by_id[export_line.model_code]
        materials = allocate_exact_bom(
            baseline,
            export_line,
            variant,
            stock_by_material,
            material_basis_stats,
            stock_by_basis,
            import_lead_days,
            max_import_age_days,
        )
        before_snapshot = summarize_product(baseline, export_line, material_state_clone(materials))
        accepted_decisions: list[dict[str, object]] = []
        iteration_no = 0
        current_snapshot = before_snapshot
        seen_states: set[tuple[object, ...]] = set()

        while True:
            if current_snapshot.passes_rvc:
                break
            state_signature = product_state_signature(materials)
            if state_signature in seen_states:
                break
            seen_states.add(state_signature)
            iteration_no += 1
            material_plans: list[MaterialPlan] = []
            for material in materials:
                context = CandidateContext(
                    starting_point_id=starting_point_id,
                    sequence_no=sequence_no,
                    iteration_no=iteration_no,
                    model_code=export_line.model_code,
                    variant_id=variant.variant_id,
                    declaration_no=export_line.declaration_no,
                )
                if collect_candidate_rows:
                    for row in build_candidate_rows(
                        baseline,
                        context,
                        export_line,
                        variant,
                        materials,
                        current_snapshot,
                        material,
                        stock_by_basis,
                        staff_substitutes,
                        import_lead_days,
                        max_import_age_days,
                        stock_by_material=stock_by_material,
                        replacement_mode=replacement_mode,
                        explicit_substitute_index=explicit_substitute_index,
                    ):
                        candidate_rows_by_key[candidate_row_key(row)] = row
                plan = build_material_plan(
                    baseline,
                    export_line,
                    variant,
                    materials,
                    current_snapshot,
                    material,
                    stock_by_basis,
                    staff_substitutes,
                    import_lead_days,
                    max_import_age_days,
                    stock_by_material=stock_by_material,
                    replacement_mode=replacement_mode,
                    explicit_substitute_index=explicit_substitute_index,
                )
                if plan is not None and plan_improves(plan, current_snapshot):
                    material_plans.append(plan)

            if not material_plans:
                break
            chosen_plan = max(material_plans, key=plan_sort_key)
            target_material = next(
                item for item in materials if item.original_material_code == chosen_plan.original_material_code
            )
            apply_material_plan(target_material, chosen_plan, bucket_by_id)
            current_snapshot = summarize_product(baseline, export_line, material_state_clone(materials))
            grouped_allocations = group_candidate_allocations(chosen_plan.allocations)
            accepted_decisions.append(
                {
                    "iteration_no": iteration_no,
                    "original_material_code": chosen_plan.original_material_code,
                    "custom_code_basis": chosen_plan.custom_code_basis,
                    "covered_qty": chosen_plan.covered_qty,
                    "remaining_unmet_qty": chosen_plan.remaining_unmet_qty,
                    "realized_cost_usd": chosen_plan.realized_cost_usd,
                    "non_origin_value_usd": chosen_plan.non_origin_value_usd,
                    "rvc_after": current_snapshot.rvc_percent,
                    "reaches_target_after_decision": current_snapshot.passes_rvc,
                    "candidate_materials": [
                        {
                            "candidate_material_code": material_code,
                            "allocated_qty": sum(item.allocated_qty for item in allocations),
                            "staff_known_substitute": any(item.staff_known_substitute for item in allocations),
                            "mapping_confidence": max(
                                (item.mapping_confidence for item in allocations),
                                key=confidence_rank,
                            ),
                            "allocated_buckets": [asdict(item) for item in allocations],
                        }
                        for material_code, allocations in grouped_allocations.items()
                    ],
                }
            )
            if current_snapshot.passes_rvc:
                break

        after_snapshot = summarize_product(baseline, export_line, material_state_clone(materials))
        if accepted_decisions:
            changed_products.append(export_line.model_code)
            changed_materials.update(
                item["original_material_code"] for item in accepted_decisions
            )

        material_mappings = []
        for material in materials:
            replacement_codes = sorted({item.material_code for item in material.replacement_allocations})
            if replacement_codes:
                material_mappings.append(
                    {
                        "original_material_code": material.original_material_code,
                        "custom_code_basis": material.custom_code_basis,
                        "replacement_material_codes": replacement_codes,
                    }
                )
        snapshot_payload = {
            "shipment_id": shipment_id,
            "starting_point_id": starting_point_id,
            "replacement_mode": replacement_mode,
            "sequence_no": sequence_no,
            "model_code": export_line.model_code,
            "variant_id": variant.variant_id,
            "bom_code": variant.bom_code,
            "declaration_no": export_line.declaration_no,
            "export_qty": export_line.quantity,
            "export_date": export_line.export_date,
            "rvc_before": before_snapshot.rvc_percent,
            "rvc_after": after_snapshot.rvc_percent,
            "bom_before": snapshot_to_payload(before_snapshot),
            "accepted_replacement_decisions": accepted_decisions,
            "bom_after": snapshot_to_payload(after_snapshot),
            "material_mappings": material_mappings,
            "local_stock_state_after_product": stock_state_summary(stock_by_material),
        }
        consumption_rows.extend(
            product_consumption_rows(
                starting_point_id,
                sequence_no,
                export_line,
                variant,
                materials,
            )
        )
        after_product_stock_rows.extend(
            stock_after_product_rows(
                starting_point_id,
                sequence_no,
                export_line,
                snapshot_payload["local_stock_state_after_product"],
            )
        )
        snapshot_path = snapshot_dir / f"{sequence_no:02d}-{export_line.model_code}.json"
        snapshot_path.write_text(
            json.dumps(snapshot_payload, indent=2, ensure_ascii=False, default=str) + "\n",
            encoding="utf-8",
        )

        final_products.append(
            {
                "sequence_no": sequence_no,
                "model_code": export_line.model_code,
                "variant_id": variant.variant_id,
                "bom_code": variant.bom_code,
                "declaration_no": export_line.declaration_no,
                "rvc_before": before_snapshot.rvc_percent,
                "rvc_after": after_snapshot.rvc_percent,
                "stock_before": before_snapshot.stock_sufficient,
                "stock_after": after_snapshot.stock_sufficient,
                "passes_before": before_snapshot.passes_rvc,
                "passes_after": after_snapshot.passes_rvc,
                "unmet_before": before_snapshot.unmet_qty_total,
                "unmet_after": after_snapshot.unmet_qty_total,
                "changed_materials": ";".join(
                    sorted({item["original_material_code"] for item in accepted_decisions})
                ),
                "snapshot_path": str(snapshot_path),
            }
        )
        product_status_rows.append(
            {
                "starting_point_id": starting_point_id,
                **final_products[-1],
            }
        )

    candidate_rows = sorted(
        candidate_rows_by_key.values(),
        key=lambda row: (
            row["sequence_no"],
            row["model_code"],
            row["original_material_code"],
            row["candidate_rank"],
            row["candidate_material_code"],
        ),
    )
    margins = [item["rvc_after"] - baseline.TARGET_RVC for item in final_products if item["rvc_after"] is not None]
    seed_summary = {
        "shipment_id": shipment_id,
        "starting_point_id": starting_point_id,
        "pass_count": sum(1 for item in final_products if item["passes_after"]),
        "final_unmet_qty": sum(item["unmet_after"] for item in final_products),
        "min_margin": min(margins) if margins else None,
        "changed_products": changed_products,
        "changed_materials": sorted(changed_materials),
        "product_status": final_products,
        "final_bom_state_after_full_run": final_products,
        "metadata": {
            "replacement_mode": replacement_mode,
            "future_runtime_selectable": [
                "bom_selection_strategy",
                "sequence_strategy",
                "replacement_pass_scope",
                "replacement_objective",
            ],
            "hard_coded_defaults": {
                "starting_points": [],
                "pass_1_stock_scope": "same_custom_code_clean_plus_ambiguity_plus_no_lookup",
                "objective": "reach_stock_sufficient_and_rvc_35_then_stop_per_product",
                "fallback": "maximize_pass_feasibility_then_reduce_non_origin_cost",
                "bom_output": "per_product_bom_snapshot",
            },
        },
    }
    return {
        "seed_summary": seed_summary,
        "candidate_rows": candidate_rows,
        "product_status_rows": product_status_rows,
        "stock_exports": {
            "initial_bucket_rows": initial_bucket_rows,
            "consumption_rows": consumption_rows,
            "after_product_stock_rows": after_product_stock_rows,
            "final_bucket_rows": stock_bucket_rows(stock_by_material),
            "final_material_rows": stock_material_summary_rows(stock_by_material),
        },
    }


def write_csv(path: Path, rows: list[dict[str, object]], fieldnames: list[str]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            writer.writerow({field: row.get(field) for field in fieldnames})


def write_sheet(
    workbook: openpyxl.Workbook,
    title: str,
    rows: list[dict[str, object]],
    fieldnames: list[str],
) -> None:
    if getattr(workbook, "write_only", False):
        sheet = workbook.create_sheet(title[:31])
        sheet.append(fieldnames)
        for row in rows:
            sheet.append([row.get(field) for field in fieldnames])
        return
    if workbook.worksheets and workbook.active.max_row == 1 and workbook.active.max_column == 1 and workbook.active["A1"].value is None:
        sheet = workbook.active
        sheet.title = title[:31]
        sheet.delete_rows(1, 1)
    else:
        sheet = workbook.create_sheet(title[:31])
    sheet.append(fieldnames)
    for row in rows:
        sheet.append([row.get(field) for field in fieldnames])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def write_review_workbook(
    path: Path,
    basis_rows: list[dict[str, object]],
    seed_summaries: list[dict[str, object]],
    product_status_rows: list[dict[str, object]],
    candidate_rows_by_seed: dict[str, list[dict[str, object]]],
) -> None:
    workbook = openpyxl.Workbook(write_only=True)
    write_sheet(
        workbook,
        "Seed Summary",
        [
            {
                "starting_point_id": item["starting_point_id"],
                "pass_count": item["pass_count"],
                "final_unmet_qty": item["final_unmet_qty"],
                "min_margin": item["min_margin"],
                "changed_products": ";".join(item["changed_products"]),
                "changed_materials": ";".join(item["changed_materials"]),
            }
            for item in seed_summaries
        ],
        [
            "starting_point_id",
            "pass_count",
            "final_unmet_qty",
            "min_margin",
            "changed_products",
            "changed_materials",
        ],
    )
    write_sheet(
        workbook,
        "Product Status",
        product_status_rows,
        [
            "starting_point_id",
            "sequence_no",
            "model_code",
            "variant_id",
            "bom_code",
            "declaration_no",
            "rvc_before",
            "rvc_after",
            "stock_before",
            "stock_after",
            "passes_before",
            "passes_after",
            "unmet_before",
            "unmet_after",
            "changed_materials",
            "snapshot_path",
        ],
    )
    write_sheet(
        workbook,
        "Custom Code Basis",
        basis_rows,
        [
            "custom_code_basis",
            "erp_material_code",
            "candidate_material_code",
            "bucket_id",
            "tracking_key",
            "candidate_class",
            "mapping_status",
            "mapping_confidence",
            "evidence_source",
            "remaining_qty",
            "unit_price_usd",
            "origin",
            "shipment_variant_scope",
            "variant_scope",
            "variant_hit_count",
            "source",
            "source_row_no",
            "declaration_no",
            "declaration_item_no",
            "import_date",
            "hs_code",
            "name",
        ],
    )
    for seed, rows in candidate_rows_by_seed.items():
        write_sheet(
            workbook,
            f"Cand {seed}",
            rows,
            [
                "starting_point_id",
                "sequence_no",
                "iteration_no",
                "model_code",
                "variant_id",
                "declaration_no",
                "original_material_code",
                "custom_code_basis",
                "candidate_rank",
                "candidate_material_code",
                "admissible_qty",
                "allocated_preview_qty",
                "clean_qty_available",
                "date_blocked_qty_available",
                "ambiguity_qty_available",
                "ambiguity_date_blocked_qty_available",
                "review_scope",
                "risk_flags",
                "risk_flag_count",
                "commit_scope_eligible",
                "estimate_scope",
                "mapping_confidence",
                "confidence_rank",
                "mapping_evidence_source",
                "staff_known_substitute",
                "estimated_rvc_after",
                "estimated_rvc_gain",
                "would_reach_target",
                "realized_cost_usd",
                "variant_hit_scope",
                "candidate_class",
                "remaining_unmet_after_candidate",
                "bucket_cost_basis",
            ],
        )
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def build_basis_csv_rows(
    basis_buckets: list[ReplacementBasisBucket],
) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for bucket in basis_buckets:
        rows.append(
            {
                "custom_code_basis": bucket.custom_code_basis,
                "erp_material_code": bucket.erp_material_code,
                "candidate_material_code": bucket.candidate_material_code,
                "bucket_id": bucket.bucket_id,
                "tracking_key": bucket.tracking_key,
                "candidate_class": bucket.candidate_class,
                "mapping_status": bucket.mapping_status,
                "mapping_confidence": bucket.mapping_confidence,
                "evidence_source": bucket.evidence_source,
                "remaining_qty": bucket.remaining_qty,
                "unit_price_usd": bucket.unit_price_usd,
                "origin": bucket.origin,
                "shipment_variant_scope": bucket.shipment_variant_scope,
                "variant_scope": ";".join(bucket.variant_scope),
                "variant_hit_count": bucket.variant_hit_count,
                "source": bucket.source,
                "source_row_no": bucket.source_row_no,
                "declaration_no": bucket.declaration_no,
                "declaration_item_no": bucket.declaration_item_no,
                "import_date": bucket.import_date,
                "hs_code": bucket.hs_code,
                "name": bucket.name,
            }
        )
    return rows


def resolve_starting_point_export_lines(
    shipment_payload: dict[str, object],
    shipment_id: str,
    export_rows: dict[str, list[object]],
) -> dict[str, list[object]]:
    export_lookup = {
        f"{line.model_code}:{line.declaration_no}:{line.declaration_item_no}": line
        for line in export_rows[shipment_id]
    }
    invoice_lookup = {
        f"{line.model_code}:{line.invoice_no or 'no-invoice'}:{line.declaration_no}:{line.declaration_item_no}": line
        for line in export_rows[shipment_id]
    }
    resolved: dict[str, list[object]] = {}
    for starting_point_id, item in shipment_payload.items():
        lines: list[object] = []
        for key in item["export_order"]:
            line = export_lookup.get(key) or invoice_lookup.get(key)
            if line is None:
                raise SystemExit(f"Unable to resolve export line {key} for {starting_point_id}")
            lines.append(line)
        resolved[starting_point_id] = lines
    return resolved


def load_starting_point_export_lines(
    baseline_starting_points_path: Path,
    shipment_id: str,
    export_rows: dict[str, list[object]],
) -> dict[str, list[object]]:
    payload = json.loads(baseline_starting_points_path.read_text(encoding="utf-8"))
    return resolve_starting_point_export_lines(payload[shipment_id], shipment_id, export_rows)


def bucket_non_origin_unit_value(baseline: object, bucket: ReplacementBasisBucket) -> float:
    if baseline.is_vietnam_origin(bucket.origin):
        return 0.0
    return bucket.unit_price_usd


def analyze_material_replaceability(
    baseline: object,
    export_line: object,
    variant_id: str,
    material_code: str,
    current_non_origin_unit_value: float,
    stock_by_basis: dict[str, list[ReplacementBasisBucket]],
    material_basis_stats: dict[str, dict[str, BasisStats]],
    import_lead_days: int,
    max_import_age_days: int,
) -> dict[str, object]:
    custom_code_basis = choose_custom_code_basis(material_code, material_basis_stats, stock_by_basis)
    alt_codes: set[str] = set()
    better_alt_codes: set[str] = set()
    alt_qty = 0.0
    better_alt_qty = 0.0
    for bucket in stock_by_basis.get(custom_code_basis, []):
        if not bucket_is_eligible(
            baseline,
            bucket,
            export_line,
            variant_id,
            import_lead_days,
            max_import_age_days,
        ):
            continue
        candidate_code = bucket.candidate_material_code
        if not candidate_code or candidate_code == material_code:
            continue
        alt_codes.add(candidate_code)
        alt_qty += bucket.remaining_qty
        if bucket_non_origin_unit_value(baseline, bucket) + COST_EPS < current_non_origin_unit_value:
            better_alt_codes.add(candidate_code)
            better_alt_qty += bucket.remaining_qty
    return {
        "custom_code_basis": custom_code_basis,
        "alt_code_count": len(alt_codes),
        "alt_qty": alt_qty,
        "better_alt_code_count": len(better_alt_codes),
        "better_alt_qty": better_alt_qty,
    }


def score_replaceability_scenario(
    baseline: object,
    scenario: dict[str, object],
    export_lookup: dict[tuple[str, str], object],
    stock_by_basis: dict[str, list[ReplacementBasisBucket]],
    material_basis_stats: dict[str, dict[str, BasisStats]],
    import_lead_days: int,
    max_import_age_days: int,
) -> dict[str, object]:
    replaceable_material_count = 0
    replaceable_candidate_codes = 0
    unmet_replaceable_material_count = 0
    unmet_replaceable_qty = 0.0
    cost_down_material_count = 0
    cost_down_non_origin_value_usd = 0.0
    better_candidate_codes = 0
    for product_result in scenario["product_results"]:
        export_line = export_lookup[(product_result["model_code"], product_result["declaration_no"])]
        variant_id = clean_text(product_result["variant_id"])
        for material in product_result["materials"]:
            material_code = clean_text(material["material_code"])
            need_qty = float(material["need_qty"])
            unmet_qty = float(material["unmet_qty"])
            current_non_origin_value = float(material["non_origin_value_usd"])
            current_non_origin_unit_value = current_non_origin_value / need_qty if need_qty > EPS else 0.0
            metrics = analyze_material_replaceability(
                baseline,
                export_line,
                variant_id,
                material_code,
                current_non_origin_unit_value,
                stock_by_basis,
                material_basis_stats,
                import_lead_days,
                max_import_age_days,
            )
            if metrics["alt_code_count"] > 0:
                replaceable_material_count += 1
                replaceable_candidate_codes += int(metrics["alt_code_count"])
            if unmet_qty > EPS and metrics["alt_qty"] > EPS:
                unmet_replaceable_material_count += 1
                unmet_replaceable_qty += min(unmet_qty, float(metrics["alt_qty"]))
            if current_non_origin_value > EPS and metrics["better_alt_code_count"] > 0:
                cost_down_material_count += 1
                cost_down_non_origin_value_usd += current_non_origin_value
                better_candidate_codes += int(metrics["better_alt_code_count"])
    return {
        "scenario_id": scenario["scenario_id"],
        "stock_sufficient": scenario["stock_sufficient"],
        "passes_rvc": scenario["passes_rvc"],
        "min_margin": scenario["min_margin"],
        "total_unmet_qty": scenario["total_unmet_qty"],
        "replaceable_material_count": replaceable_material_count,
        "replaceable_candidate_codes": replaceable_candidate_codes,
        "unmet_replaceable_material_count": unmet_replaceable_material_count,
        "unmet_replaceable_qty": unmet_replaceable_qty,
        "cost_down_material_count": cost_down_material_count,
        "cost_down_non_origin_value_usd": cost_down_non_origin_value_usd,
        "better_candidate_codes": better_candidate_codes,
    }


def sort_key_replaceability_unmet(metrics: dict[str, object]) -> tuple[float, int, int, float, float, str]:
    min_margin = float(metrics["min_margin"]) if metrics["min_margin"] is not None else float("-inf")
    return (
        float(metrics["unmet_replaceable_qty"]),
        int(metrics["unmet_replaceable_material_count"]),
        int(metrics["replaceable_candidate_codes"]),
        -float(metrics["total_unmet_qty"]),
        min_margin,
        clean_text(metrics["scenario_id"]),
    )


def sort_key_replaceability_cost_down(metrics: dict[str, object]) -> tuple[float, int, int, float, float, str]:
    min_margin = float(metrics["min_margin"]) if metrics["min_margin"] is not None else float("-inf")
    return (
        float(metrics["cost_down_non_origin_value_usd"]),
        int(metrics["cost_down_material_count"]),
        int(metrics["better_candidate_codes"]),
        -float(metrics["total_unmet_qty"]),
        min_margin,
        clean_text(metrics["scenario_id"]),
    )


def make_starting_point_payload(
    starting_point_id: str,
    variant_strategy: str,
    scenario: dict[str, object],
    export_lines: list[object],
) -> dict[str, object]:
    return {
        "starting_point_id": starting_point_id,
        "variant_strategy": variant_strategy,
        "sequence_strategy": "declaration_order",
        "export_order": [
            f"{line.model_code}:{line.declaration_no}:{line.declaration_item_no}"
            for line in export_lines
        ],
        "scenario": scenario,
    }


def extend_starting_points_with_heuristics(
    baseline: object,
    shipment_id: str,
    starting_points_payload: dict[str, object],
    baseline_scenarios_path: Path,
    export_lines: list[object],
    export_lookup: dict[tuple[str, str], object],
    stock_by_basis: dict[str, list[ReplacementBasisBucket]],
    material_basis_stats: dict[str, dict[str, BasisStats]],
    import_lead_days: int,
    max_import_age_days: int,
) -> tuple[dict[str, object], list[dict[str, object]]]:
    payload = json.loads(baseline_scenarios_path.read_text(encoding="utf-8"))
    shipment_scenarios = payload.get(shipment_id, [])
    existing_scenario_ids = {
        clean_text(item["scenario"]["scenario_id"])
        for item in starting_points_payload.values()
    }
    scored: list[tuple[dict[str, object], dict[str, object]]] = []
    analysis_rows: list[dict[str, object]] = []
    for scenario in shipment_scenarios:
        metrics = score_replaceability_scenario(
            baseline,
            scenario,
            export_lookup,
            stock_by_basis,
            material_basis_stats,
            import_lead_days,
            max_import_age_days,
        )
        scored.append((scenario, metrics))
        analysis_rows.append(metrics)

    extended = dict(starting_points_payload)
    selected_scenario_ids: set[str] = set()
    selections = [
        (
            "replaceability_unmet_best",
            "highest_replaceable_unmet_qty_across_baseline_scenarios",
            sort_key_replaceability_unmet,
        ),
        (
            "replaceability_cost_down_best",
            "highest_cost_down_potential_across_baseline_scenarios",
            sort_key_replaceability_cost_down,
        ),
    ]
    for starting_point_id, variant_strategy, sort_key in selections:
        ranked = sorted(scored, key=lambda item: sort_key(item[1]), reverse=True)
        for scenario, metrics in ranked:
            scenario_id = clean_text(scenario["scenario_id"])
            if scenario_id in existing_scenario_ids or scenario_id in selected_scenario_ids:
                continue
            extended[starting_point_id] = make_starting_point_payload(
                starting_point_id,
                variant_strategy,
                scenario,
                export_lines,
            )
            selected_scenario_ids.add(scenario_id)
            metrics["selected_as"] = starting_point_id
            break
    return extended, analysis_rows


def build_material_metadata_index(shared_normalized_dir: Path) -> dict[str, dict[str, str]]:
    index: dict[str, dict[str, str]] = {}
    with (shared_normalized_dir / "co-stock-tracking-updated.csv").open(encoding="utf-8") as handle:
        for row in csv.DictReader(handle):
            for key in {
                clean_text(row.get("confirmed_lookup_code")),
                clean_text(row.get("lookup_material_code")),
            }:
                if not key or key in index:
                    continue
                index[key] = {
                    "name": clean_text(row.get("name")) or key,
                    "hs_code": clean_text(row.get("hs_code")),
                    "unit": clean_text(row.get("unit")),
                }
    return index


def format_origin_status(origins: list[str]) -> str:
    return "Không có xuất xứ"


def apply_fill_to_row(
    sheet: openpyxl.worksheet.worksheet.Worksheet,
    row_idx: int,
    start_col: int,
    end_col: int,
    fill: PatternFill,
) -> None:
    for col_idx in range(start_col, end_col + 1):
        sheet.cell(row=row_idx, column=col_idx).fill = fill


def workbook_seed_dir(results_dir: Path, starting_point_id: str) -> Path:
    return results_dir / "replacement-excel" / starting_point_id


def sheet_name_for_product(sequence_no: int, model_code: str) -> str:
    return f"{sequence_no:02d}-{model_code}"[:31]


def active_allocations_from_material(material: dict[str, object]) -> list[dict[str, object]]:
    return [
        *material.get("exact_allocations", []),
        *material.get("replacement_allocations", []),
    ]


def write_rvc_header(sheet: openpyxl.worksheet.worksheet.Worksheet, export_line: object, snapshot: dict[str, object]) -> None:
    sheet.merge_cells("B3:Q3")
    sheet.merge_cells("B4:Q4")
    sheet["B2"] = "Phụ lục VIII"
    sheet["B3"] = "BẢNG KÊ KHAI HÀNG HÓA XUẤT KHẨU ĐẠT TIÊU CHÍ “RVC”"
    sheet["B4"] = "(ban hành kèm theo Thông tư số 05/2018/TT-BCT ngày 03 tháng 4 năm 2018 quy định về xuất xứ hàng hóa)"
    sheet["B6"] = f"Tên Thương nhân: {MERCHANT_NAME}"
    sheet["B7"] = f"Mã số thuế : {MERCHANT_TAX_ID}"
    sheet["J6"] = "Tiêu chí áp dụng: RVC 35% + CTSH"
    sheet["J7"] = "Tên hàng hóa : "
    sheet["K7"] = clean_text(getattr(export_line, "name", ""))
    sheet["B8"] = (
        f"Tờ khai hải quan xuất khẩu số : {export_line.declaration_no}/E42 ngày : "
        f"{clean_text(getattr(export_line, 'export_date', ''))}"
    )
    sheet["J8"] = "Mã HS của hàng hóa (6 số ) :"
    sheet["K8"] = clean_text(getattr(export_line, "hs_code", ""))[:6]
    sheet["J9"] = "Số lượng : "
    sheet["K9"] = getattr(export_line, "quantity", "")
    sheet["M9"] = "Đơn vị tính :"
    sheet["N9"] = "Chiếc"
    sheet["J10"] = "Trị giá FOB :"
    sheet["K10"] = float(getattr(export_line, "quantity", 0.0)) * float(getattr(export_line, "unit_price_usd", 0.0))
    for cell in ("B2", "B3", "B4", "B6", "B7", "B8", "J6", "J7", "J8", "J9", "J10"):
        sheet[cell].font = Font(bold=True)
    sheet["B3"].alignment = Alignment(horizontal="center")
    sheet["B4"].alignment = Alignment(horizontal="center")


def write_rvc_table_header(sheet: openpyxl.worksheet.worksheet.Worksheet) -> None:
    headers = {
        "A12": "STT\n序号",
        "B12": "Tên nguyên phụ liệu\n原材料的名称",
        "C12": "Mã HS\nHS CODE",
        "D12": "Đơn vị tính\n计算单位",
        "E12": "Định mức sản phẩm kể cả % hao hụt\n用量已经含损耗率",
        "F12": "Đơn giá (CIF)",
        "G12": "Nhu cầu nguyên liệu sử dụng cho lô hàng\n原材料的需求用于一票货",
        "H12": "Trị giá có xuất xứ",
        "I12": "Trị giá không có xuất xứ",
        "J12": "Nước xuất xứ\n原产国",
        "K12": "Tờ khai nhập khẩu / Hóa đơn\n报关单号/发票",
        "L12": "Ngày\n报关日期",
        "M12": "Ghi chú thay thế",
    }
    for cell, value in headers.items():
        sheet[cell] = value
        sheet[cell].font = Font(bold=True)
        sheet[cell].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[12].height = 36


def append_rvc_material_rows(
    sheet: openpyxl.worksheet.worksheet.Worksheet,
    snapshot: dict[str, object],
    export_line: object,
    metadata_index: dict[str, dict[str, str]],
) -> int:
    row_idx = 13
    line_no = 1
    export_qty = float(snapshot["export_qty"])
    for material in snapshot["bom_after"]["materials"]:
        allocations = active_allocations_from_material(material)
        unmet_qty = float(material.get("unmet_qty", 0.0))
        if not allocations:
            metadata = metadata_index.get(material["original_material_code"], {})
            sheet.append(
                [
                    line_no,
                    material["original_material_code"],
                    metadata.get("hs_code", ""),
                    metadata.get("unit", ""),
                    material["need_qty"] / export_qty if export_qty else 0.0,
                    "",
                    material["need_qty"],
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    f"Thiếu NVL: chưa có allocation, thiếu {unmet_qty:.4f}",
                ]
            )
            apply_fill_to_row(sheet, row_idx, 1, 13, UNMET_FILL)
            row_idx += 1
            line_no += 1
            continue
        grouped: dict[str, list[dict[str, object]]] = defaultdict(list)
        for allocation in allocations:
            grouped[allocation["material_code"]].append(allocation)
        for material_code, bucket_allocations in grouped.items():
            metadata = metadata_index.get(material_code, {})
            total_qty = sum(float(item["allocated_qty"]) for item in bucket_allocations)
            total_value = sum(float(item["allocated_qty"]) * float(item["unit_price_usd"]) for item in bucket_allocations)
            unit_price = total_value / total_qty if total_qty else 0.0
            origins = [clean_text(item["origin"]) for item in bucket_allocations]
            replace_note = ""
            if material_code != material["original_material_code"]:
                replace_note = f"{material['original_material_code']} -> {material_code}"
            if unmet_qty > EPS:
                unmet_note = f"Thiếu NVL: {unmet_qty:.4f}"
                replace_note = f"{replace_note} | {unmet_note}" if replace_note else unmet_note
            sheet.append(
                [
                    line_no,
                    metadata.get("name", material_code),
                    metadata.get("hs_code", ""),
                    metadata.get("unit", ""),
                    total_qty / export_qty if export_qty else 0.0,
                    unit_price,
                    total_qty,
                    "",
                    total_value,
                    format_origin_status(origins),
                    "\n".join(sorted({clean_text(item["declaration_no"]) for item in bucket_allocations if clean_text(item["declaration_no"])})),
                    "\n".join(sorted({clean_text(item["import_date"]) for item in bucket_allocations if clean_text(item["import_date"])})),
                    replace_note,
                ]
            )
            if unmet_qty > EPS:
                apply_fill_to_row(sheet, row_idx, 1, 13, UNMET_FILL)
            row_idx += 1
            line_no += 1
    return row_idx


def write_rvc_footer(
    sheet: openpyxl.worksheet.worksheet.Worksheet,
    start_row: int,
    snapshot: dict[str, object],
    export_line: object,
) -> None:
    fob_value = float(getattr(export_line, "quantity", 0.0)) * float(getattr(export_line, "unit_price_usd", 0.0))
    non_origin_value = float(snapshot["bom_after"]["non_origin_value_usd"])
    rvc_after = snapshot["rvc_after"]
    row = start_row + 1
    sheet[f"B{row}"] = "Trị giá FOB"
    sheet[f"K{row}"] = fob_value
    row += 1
    sheet[f"B{row}"] = "Trị giá CIF nguyên liệu đầu vào không có xuất xứ"
    sheet[f"K{row}"] = non_origin_value
    row += 1
    sheet[f"B{row}"] = "Công thức tính RVC gián tiếp = (FOB - CIF không có xuất xứ) / FOB x 100"
    sheet[f"K{row}"] = rvc_after
    row += 1
    if rvc_after is None:
        message = "Kết luận: chưa tính được RVC"
    elif snapshot["bom_after"]["passes_rvc"]:
        message = f"Kết luận: Sản phẩm đạt tiêu chí RVC {rvc_after:.2f}% + CTSH"
    elif snapshot["bom_after"]["stock_sufficient"]:
        message = f"Kết luận: Sản phẩm chưa đạt tiêu chí RVC {rvc_after:.2f}% + CTSH"
    else:
        message = (
            f"Kết luận: Sản phẩm chưa đạt do thiếu NVL; RVC tính được là {rvc_after:.2f}% "
            f"nhưng unmet còn {snapshot['bom_after']['unmet_qty_total']:.4f}"
        )
    sheet[f"B{row}"] = message
    sheet[f"B{row}"].font = Font(bold=True)


def finalize_standard_sheet(sheet: openpyxl.worksheet.worksheet.Worksheet) -> None:
    widths = {
        "A": 8,
        "B": 60,
        "C": 12,
        "D": 12,
        "E": 14,
        "F": 14,
        "G": 14,
        "H": 16,
        "I": 16,
        "J": 18,
        "K": 20,
        "L": 14,
        "M": 24,
    }
    for col, width in widths.items():
        sheet.column_dimensions[col].width = width
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def write_rvc_standard_workbook(
    path: Path,
    snapshots: list[dict[str, object]],
    export_lookup: dict[tuple[str, str], object],
    metadata_index: dict[str, dict[str, str]],
) -> None:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for snapshot in snapshots:
        export_line = export_lookup[(snapshot["model_code"], snapshot["declaration_no"])]
        sheet = workbook.create_sheet(sheet_name_for_product(snapshot["sequence_no"], snapshot["model_code"]))
        write_rvc_header(sheet, export_line, snapshot)
        write_rvc_table_header(sheet)
        next_row = append_rvc_material_rows(sheet, snapshot, export_line, metadata_index)
        write_rvc_footer(sheet, next_row, snapshot, export_line)
        finalize_standard_sheet(sheet)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def write_bom_before_after_workbook(path: Path, snapshots: list[dict[str, object]]) -> None:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for snapshot in snapshots:
        sheet = workbook.create_sheet(sheet_name_for_product(snapshot["sequence_no"], snapshot["model_code"]))
        sheet["A1"] = f"Seed: {snapshot['starting_point_id']}"
        sheet["A2"] = f"Model: {snapshot['model_code']}"
        sheet["A3"] = f"Variant: {snapshot['variant_id']}"
        sheet["A4"] = f"RVC before: {snapshot['rvc_before']}"
        sheet["A5"] = f"RVC after: {snapshot['rvc_after']}"
        headers = [
            "Material",
            "Basis Before",
            "Need Before",
            "Active Before",
            "Unmet Before",
            "Basis After",
            "Need After",
            "Active After",
            "Unmet After",
            "Changed",
        ]
        for idx, value in enumerate(headers, start=1):
            cell = sheet.cell(row=7, column=idx, value=value)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
        before_index = {
            item["original_material_code"]: item
            for item in snapshot["bom_before"]["materials"]
        }
        after_index = {
            item["original_material_code"]: item
            for item in snapshot["bom_after"]["materials"]
        }
        row_idx = 8
        for material_code in sorted(set(before_index) | set(after_index)):
            before = before_index.get(material_code, {})
            after = after_index.get(material_code, {})
            before_active = sorted({item["material_code"] for item in active_allocations_from_material(before)})
            after_active = sorted({item["material_code"] for item in active_allocations_from_material(after)})
            changed = before_active != after_active or round(float(before.get("unmet_qty", 0.0)), 9) != round(float(after.get("unmet_qty", 0.0)), 9)
            values = [
                material_code,
                before.get("custom_code_basis", ""),
                before.get("need_qty", ""),
                "; ".join(before_active),
                before.get("unmet_qty", ""),
                after.get("custom_code_basis", ""),
                after.get("need_qty", ""),
                "; ".join(after_active),
                after.get("unmet_qty", ""),
                "yes" if changed else "",
            ]
            for col_idx, value in enumerate(values, start=1):
                sheet.cell(row=row_idx, column=col_idx, value=value)
            if changed:
                apply_fill_to_row(sheet, row_idx, 1, len(headers), CHANGED_FILL)
            if float(before.get("unmet_qty", 0.0)) > EPS:
                sheet.cell(row=row_idx, column=5).fill = UNMET_FILL
            if float(after.get("unmet_qty", 0.0)) > EPS:
                sheet.cell(row=row_idx, column=9).fill = UNMET_FILL
            row_idx += 1
        for col, width in zip("ABCDEFGHIJ", [18, 16, 14, 36, 14, 16, 14, 36, 14, 10]):
            sheet.column_dimensions[col].width = width
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def write_stock_tracking_workbook(
    path: Path,
    initial_bucket_rows: list[dict[str, object]],
    consumption_rows: list[dict[str, object]],
    after_product_rows: list[dict[str, object]],
    final_bucket_rows: list[dict[str, object]],
    final_material_rows: list[dict[str, object]],
) -> None:
    workbook = openpyxl.Workbook(write_only=True)
    write_sheet(
        workbook,
        "Initial Buckets",
        initial_bucket_rows,
        [
            "material_code",
            "bucket_material_code",
            "candidate_material_code",
            "custom_code_basis",
            "bucket_id",
            "tracking_key",
            "declaration_no",
            "declaration_item_no",
            "import_date",
            "source",
            "source_row_no",
            "hs_code",
            "name",
            "origin",
            "unit_price_usd",
            "exchange_rate",
            "remaining_qty",
            "candidate_class",
            "mapping_status",
            "mapping_confidence",
            "evidence_source",
            "shipment_variant_scope",
            "variant_scope",
            "variant_hit_count",
        ],
    )
    write_sheet(
        workbook,
        "Consumption Ledger",
        consumption_rows,
        [
            "starting_point_id",
            "sequence_no",
            "model_code",
            "variant_id",
            "bom_code",
            "declaration_no",
            "original_material_code",
            "allocation_type",
            "allocated_material_code",
            "custom_code_basis",
            "bucket_id",
            "tracking_key",
            "bucket_declaration_no",
            "bucket_declaration_item_no",
            "import_date",
            "allocated_qty",
            "unit_price_usd",
            "allocated_value_usd",
            "origin",
            "source",
            "source_row_no",
            "candidate_class",
            "mapping_status",
            "mapping_confidence",
            "evidence_source",
            "staff_known_substitute",
        ],
    )
    write_sheet(
        workbook,
        "Stock After Product",
        after_product_rows,
        [
            "starting_point_id",
            "sequence_no",
            "model_code",
            "declaration_no",
            "material_code",
            "remaining_qty",
            "bucket_count",
        ],
    )
    write_sheet(
        workbook,
        "Final Buckets",
        final_bucket_rows,
        [
            "material_code",
            "bucket_material_code",
            "candidate_material_code",
            "custom_code_basis",
            "bucket_id",
            "tracking_key",
            "declaration_no",
            "declaration_item_no",
            "import_date",
            "source",
            "source_row_no",
            "hs_code",
            "name",
            "origin",
            "unit_price_usd",
            "exchange_rate",
            "remaining_qty",
            "candidate_class",
            "mapping_status",
            "mapping_confidence",
            "evidence_source",
            "shipment_variant_scope",
            "variant_scope",
            "variant_hit_count",
        ],
    )
    write_sheet(
        workbook,
        "Final By Material",
        final_material_rows,
        [
            "material_code",
            "remaining_qty",
            "bucket_count",
        ],
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def load_seed_snapshots(results_dir: Path, starting_point_id: str) -> list[dict[str, object]]:
    snapshot_dir = results_dir / "replacement-bom-snapshots" / starting_point_id
    snapshots: list[dict[str, object]] = []
    for path in sorted(snapshot_dir.glob("*.json")):
        snapshots.append(json.loads(path.read_text(encoding="utf-8")))
    return snapshots


def write_summary_doc(
    path: Path,
    shipment_id: str,
    policy: object,
    seed_summaries: list[dict[str, object]],
    starting_point_ids: list[str],
) -> None:
    lines: list[str] = []
    lines.append("# Growatt B282 Replacement Runner")
    lines.append("")
    lines.append(f"- Shipment: `{shipment_id}`")
    lines.append(f"- Policy version: `{policy.policy_version}`")
    lines.append(f"- Starting points: `{', '.join(starting_point_ids)}`")
    lines.append("- Commit scope: `same custom code + clean + ambiguity + no_lookup_code`")
    lines.append("- Review scope: `same custom code + clean + ambiguity + no_lookup_code + date-blocked`")
    lines.append("- Objective: reach `stock sufficient + RVC >= 35`, then stop replacement for that product")
    lines.append("- Fallback: maximize pass feasibility, then reduce `non-origin cost`")
    lines.append("")
    lines.append("| Starting Point | Pass Count | Final Unmet Qty | Min Margin | Changed Products | Changed Materials |")
    lines.append("| --- | ---: | ---: | ---: | --- | --- |")
    for summary in seed_summaries:
        min_margin = "n/a" if summary["min_margin"] is None else f"{summary['min_margin']:.2f}"
        lines.append(
            f"| {summary['starting_point_id']} | {summary['pass_count']} | "
            f"{summary['final_unmet_qty']:.4f} | "
            f"{min_margin} | "
            f"{';'.join(summary['changed_products']) or 'none'} | "
            f"{';'.join(summary['changed_materials']) or 'none'} |"
        )
    lines.append("")
    for summary in seed_summaries:
        lines.append(f"## {summary['starting_point_id']}")
        lines.append("")
        lines.append(
            f"- Pass count: `{summary['pass_count']}` | final unmet qty: `{summary['final_unmet_qty']:.4f}` | "
            f"min margin: `{summary['min_margin']:.2f}`" if summary["min_margin"] is not None else
            f"- Pass count: `{summary['pass_count']}` | final unmet qty: `{summary['final_unmet_qty']:.4f}` | min margin: `n/a`"
        )
        lines.append(f"- Changed products: `{';'.join(summary['changed_products']) or 'none'}`")
        lines.append(f"- Changed materials: `{';'.join(summary['changed_materials']) or 'none'}`")
        lines.append("")
        lines.append("| Seq | Model | Variant | Stock Before | Stock After | Pass Before | Pass After | RVC Before | RVC After | Unmet After |")
        lines.append("| --- | --- | --- | --- | --- | --- | --- | ---: | ---: | ---: |")
        for product in summary["product_status"]:
            rvc_before = "n/a" if product["rvc_before"] is None else f"{product['rvc_before']:.2f}"
            rvc_after = "n/a" if product["rvc_after"] is None else f"{product['rvc_after']:.2f}"
            lines.append(
                f"| {product['sequence_no']} | {product['model_code']} | {product['variant_id']} | "
                f"{'yes' if product['stock_before'] else 'no'} | {'yes' if product['stock_after'] else 'no'} | "
                f"{'yes' if product['passes_before'] else 'no'} | {'yes' if product['passes_after'] else 'no'} | "
                f"{rvc_before} | {rvc_after} | {product['unmet_after']:.4f} |"
            )
        lines.append("")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--case-dir", type=Path, default=DEFAULT_CASE_DIR)
    parser.add_argument("--shipment", default=DEFAULT_SHIPMENT)
    parser.add_argument("--workspace-dir", type=Path)
    parser.add_argument("--output-dir", type=Path)
    parser.add_argument(
        "--replacement-mode",
        choices=replacement_mode_choices(),
        default=REPLACEMENT_MODE_HEURISTIC_CUSTOM_BASIS,
    )
    parser.add_argument("--substitute-reference-path", type=Path)
    parser.add_argument("--skip-candidate-export", action="store_true")
    parser.add_argument("--output-doc", type=Path, default=DEFAULT_DOC)
    return parser.parse_args()


def run(args: argparse.Namespace) -> dict[str, object]:
    baseline = load_baseline_module()
    policy = resolve_shipment_policy(args.case_dir, args.shipment)
    shared_normalized_dir = args.case_dir / "shared" / "normalized"
    substitute_reference_path = args.substitute_reference_path or (
        technical_priority_artifact_dir(args.case_dir)
        / "substitute-reference"
        / "explicit-substitute-links.csv"
    )
    shipment_slug_value = shipment_slug(args.shipment)
    shipment_dir = args.workspace_dir or (args.case_dir / shipment_slug_value)
    workspace_results_dir = shipment_dir / "results"
    results_dir = args.output_dir or workspace_results_dir
    results_dir.mkdir(parents=True, exist_ok=True)

    variants_by_model = load_workspace_variants(shipment_dir / "normalized")
    export_rows = baseline.load_export_rows(shared_normalized_dir, {args.shipment})
    if args.shipment not in export_rows:
        raise SystemExit(f"Shipment not found in normalized exports: {args.shipment}")
    export_lookup = {
        (line.model_code, line.declaration_no): line
        for line in export_rows[args.shipment]
    }
    material_metadata_index = build_material_metadata_index(shared_normalized_dir)

    variant_by_id: dict[str, object] = {}
    for items in variants_by_model.values():
        for item in items:
            variant_by_id[item.variant_id] = item

    basis_buckets = build_replacement_basis(shared_normalized_dir, shipment_dir, shipment_slug_value)
    stock_by_material, stock_by_basis, bucket_by_id = clone_stock_indexes(basis_buckets)
    material_basis_stats = build_material_basis_stats(basis_buckets)
    staff_substitutes = read_staff_substitutes()
    explicit_substitute_index = load_explicit_substitute_index(substitute_reference_path)

    baseline_starting_points_path = workspace_results_dir / "baseline-starting-points.json"
    starting_points_payload = json.loads(baseline_starting_points_path.read_text(encoding="utf-8"))[args.shipment]
    starting_points_payload, starting_point_analysis = extend_starting_points_with_heuristics(
        baseline,
        args.shipment,
        starting_points_payload,
        workspace_results_dir / "baseline-scenarios.json",
        export_rows[args.shipment],
        export_lookup,
        stock_by_basis,
        material_basis_stats,
        policy.import_lead_days,
        policy.max_import_age_days,
    )
    replacement_starting_points_path = results_dir / "replacement-starting-points.json"
    replacement_starting_points_path.write_text(
        json.dumps({args.shipment: starting_points_payload}, indent=2, ensure_ascii=False, default=str) + "\n",
        encoding="utf-8",
    )
    starting_point_exports = resolve_starting_point_export_lines(starting_points_payload, args.shipment, export_rows)
    start_variant_by_seed = {
        starting_point_id: {
            item["model_code"]: variant_by_id[item["variant_id"]]
            for item in payload["scenario"]["product_results"]
        }
        for starting_point_id, payload in starting_points_payload.items()
    }

    basis_csv_rows = build_basis_csv_rows(basis_buckets)
    write_csv(
        results_dir / "replacement-custom-code-basis.csv",
        basis_csv_rows,
        [
            "custom_code_basis",
            "erp_material_code",
            "candidate_material_code",
            "bucket_id",
            "tracking_key",
            "candidate_class",
            "mapping_status",
            "mapping_confidence",
            "evidence_source",
            "remaining_qty",
            "unit_price_usd",
            "origin",
            "shipment_variant_scope",
            "variant_scope",
            "variant_hit_count",
            "source",
            "source_row_no",
            "declaration_no",
            "declaration_item_no",
            "import_date",
            "hs_code",
            "name",
        ],
    )
    write_csv(
        results_dir / "replacement-starting-point-analysis.csv",
        starting_point_analysis,
        [
            "scenario_id",
            "stock_sufficient",
            "passes_rvc",
            "min_margin",
            "total_unmet_qty",
            "replaceable_material_count",
            "replaceable_candidate_codes",
            "unmet_replaceable_material_count",
            "unmet_replaceable_qty",
            "cost_down_material_count",
            "cost_down_non_origin_value_usd",
            "better_candidate_codes",
            "selected_as",
        ],
    )

    seed_results: list[dict[str, object]] = []
    candidate_csvs: dict[str, Path] = {}
    candidate_rows_by_seed: dict[str, list[dict[str, object]]] = {}
    stock_exports_by_seed: dict[str, dict[str, list[dict[str, object]]]] = {}
    all_product_status_rows: list[dict[str, object]] = []
    for starting_point_id, export_lines in starting_point_exports.items():
        variant_by_model = start_variant_by_seed[starting_point_id]

        class VariantLookup(dict):
            def __missing__(self, key: str) -> object:
                raise KeyError(key)

        result = run_starting_point(
            baseline,
            shipment_id=args.shipment,
            starting_point_id=starting_point_id,
            export_lines=export_lines,
            variant_by_id=VariantLookup(variant_by_model),
            initial_stock_by_material=stock_by_material,
            material_basis_stats=material_basis_stats,
            staff_substitutes=staff_substitutes,
            import_lead_days=policy.import_lead_days,
            max_import_age_days=policy.max_import_age_days,
            output_dir=results_dir,
            replacement_mode=args.replacement_mode,
            explicit_substitute_index=explicit_substitute_index,
            collect_candidate_rows=not args.skip_candidate_export,
        )
        seed_results.append(result["seed_summary"])
        candidate_rows_by_seed[starting_point_id] = result["candidate_rows"]
        stock_exports_by_seed[starting_point_id] = result["stock_exports"]
        all_product_status_rows.extend(result["product_status_rows"])
        candidate_path = results_dir / f"replacement-candidates-{starting_point_id}.csv"
        candidate_csvs[starting_point_id] = candidate_path
        write_csv(
            candidate_path,
            result["candidate_rows"],
            [
                "starting_point_id",
                "sequence_no",
                "iteration_no",
                "model_code",
                "variant_id",
                "declaration_no",
                "original_material_code",
                "custom_code_basis",
                "candidate_rank",
                "candidate_material_code",
                "admissible_qty",
                "allocated_preview_qty",
                "clean_qty_available",
                "date_blocked_qty_available",
                "ambiguity_qty_available",
                "ambiguity_date_blocked_qty_available",
                "review_scope",
                "risk_flags",
                "risk_flag_count",
                "commit_scope_eligible",
                "estimate_scope",
                "bucket_cost_basis",
                "mapping_confidence",
                "confidence_rank",
                "mapping_evidence_source",
                "staff_known_substitute",
                "estimated_rvc_after",
                "estimated_rvc_gain",
                "would_reach_target",
                "realized_cost_usd",
                "variant_hit_scope",
                "candidate_class",
                "remaining_unmet_after_candidate",
            ],
        )

    starting_point_ids = list(starting_point_exports.keys())
    for summary in seed_results:
        summary["metadata"]["hard_coded_defaults"]["starting_points"] = starting_point_ids

    workbook_path = results_dir / "replacement-review.xlsx"
    write_review_workbook(
        workbook_path,
        basis_rows=basis_csv_rows,
        seed_summaries=seed_results,
        product_status_rows=all_product_status_rows,
        candidate_rows_by_seed=candidate_rows_by_seed,
    )
    workbook_exports: dict[str, dict[str, str]] = {}
    for starting_point_id in start_variant_by_seed:
        snapshots = load_seed_snapshots(results_dir, starting_point_id)
        seed_dir = workbook_seed_dir(results_dir, starting_point_id)
        rvc_path = seed_dir / "replacement-rvc-standard.xlsx"
        bom_path = seed_dir / "replacement-bom-before-after.xlsx"
        stock_workbook_path = seed_dir / "replacement-stock-tracking.xlsx"
        write_rvc_standard_workbook(
            rvc_path,
            snapshots,
            export_lookup,
            material_metadata_index,
        )
        write_bom_before_after_workbook(
            bom_path,
            snapshots,
        )
        stock_exports = stock_exports_by_seed[starting_point_id]
        write_stock_tracking_workbook(
            stock_workbook_path,
            stock_exports["initial_bucket_rows"],
            stock_exports["consumption_rows"],
            stock_exports["after_product_stock_rows"],
            stock_exports["final_bucket_rows"],
            stock_exports["final_material_rows"],
        )
        write_csv(
            seed_dir / "stock-initial-buckets.csv",
            stock_exports["initial_bucket_rows"],
            [
                "material_code",
                "bucket_material_code",
                "candidate_material_code",
                "custom_code_basis",
                "bucket_id",
                "tracking_key",
                "declaration_no",
                "declaration_item_no",
                "import_date",
                "source",
                "source_row_no",
                "hs_code",
                "name",
                "origin",
                "unit_price_usd",
                "exchange_rate",
                "remaining_qty",
                "candidate_class",
                "mapping_status",
                "mapping_confidence",
                "evidence_source",
                "shipment_variant_scope",
                "variant_scope",
                "variant_hit_count",
            ],
        )
        write_csv(
            seed_dir / "stock-consumption-ledger.csv",
            stock_exports["consumption_rows"],
            [
                "starting_point_id",
                "sequence_no",
                "model_code",
                "variant_id",
                "bom_code",
                "declaration_no",
                "original_material_code",
                "allocation_type",
                "allocated_material_code",
                "custom_code_basis",
                "bucket_id",
                "tracking_key",
                "bucket_declaration_no",
                "bucket_declaration_item_no",
                "import_date",
                "allocated_qty",
                "unit_price_usd",
                "allocated_value_usd",
                "origin",
                "source",
                "source_row_no",
                "candidate_class",
                "mapping_status",
                "mapping_confidence",
                "evidence_source",
                "staff_known_substitute",
            ],
        )
        write_csv(
            seed_dir / "stock-after-product.csv",
            stock_exports["after_product_stock_rows"],
            [
                "starting_point_id",
                "sequence_no",
                "model_code",
                "declaration_no",
                "material_code",
                "remaining_qty",
                "bucket_count",
            ],
        )
        write_csv(
            seed_dir / "stock-final-buckets.csv",
            stock_exports["final_bucket_rows"],
            [
                "material_code",
                "bucket_material_code",
                "candidate_material_code",
                "custom_code_basis",
                "bucket_id",
                "tracking_key",
                "declaration_no",
                "declaration_item_no",
                "import_date",
                "source",
                "source_row_no",
                "hs_code",
                "name",
                "origin",
                "unit_price_usd",
                "exchange_rate",
                "remaining_qty",
                "candidate_class",
                "mapping_status",
                "mapping_confidence",
                "evidence_source",
                "shipment_variant_scope",
                "variant_scope",
                "variant_hit_count",
            ],
        )
        write_csv(
            seed_dir / "stock-final-by-material.csv",
            stock_exports["final_material_rows"],
            [
                "material_code",
                "remaining_qty",
                "bucket_count",
            ],
        )
        workbook_exports[starting_point_id] = {
            "seed_dir": str(seed_dir),
            "rvc_workbook": str(rvc_path),
            "bom_before_after_workbook": str(bom_path),
            "stock_tracking_workbook": str(stock_workbook_path),
        }

    write_csv(
        results_dir / "replacement-product-status.csv",
        all_product_status_rows,
        [
            "starting_point_id",
            "sequence_no",
            "model_code",
            "variant_id",
            "bom_code",
            "declaration_no",
            "rvc_before",
            "rvc_after",
            "stock_before",
            "stock_after",
            "passes_before",
            "passes_after",
            "unmet_before",
            "unmet_after",
            "changed_materials",
            "snapshot_path",
        ],
    )
    write_csv(
        results_dir / "replacement-seed-summary.csv",
        [
            {
                "starting_point_id": summary["starting_point_id"],
                "pass_count": summary["pass_count"],
                "final_unmet_qty": summary["final_unmet_qty"],
                "min_margin": summary["min_margin"],
                "changed_products": ";".join(summary["changed_products"]),
                "changed_materials": ";".join(summary["changed_materials"]),
            }
            for summary in seed_results
        ],
        [
            "starting_point_id",
            "pass_count",
            "final_unmet_qty",
            "min_margin",
            "changed_products",
            "changed_materials",
        ],
    )
    (results_dir / "replacement-final-summaries.json").write_text(
        json.dumps(
            {
                "shipment_id": args.shipment,
                "policy_version": policy.policy_version,
                "workspace_dir": str(shipment_dir),
                "output_dir": str(results_dir),
                "replacement_mode": args.replacement_mode,
                "skip_candidate_export": args.skip_candidate_export,
                "substitute_reference_path": str(substitute_reference_path),
                "seed_summaries": seed_results,
                "candidate_tables": {key: str(value) for key, value in candidate_csvs.items()},
                "review_workbook": str(workbook_path),
                "workbook_exports": workbook_exports,
                "starting_points": starting_point_ids,
            },
            indent=2,
            ensure_ascii=False,
            default=str,
        )
        + "\n",
        encoding="utf-8",
    )
    (results_dir / "replacement-run-config.json").write_text(
        json.dumps(
            {
                "stage": "replacement_runner",
                "shipment_id": args.shipment,
                "policy_version": policy.policy_version,
                "import_lead_days": policy.import_lead_days,
                "max_import_age_days": policy.max_import_age_days,
                "valuation_mode": policy.valuation_mode,
                "workspace_dir": str(shipment_dir),
                "output_dir": str(results_dir),
                "replacement_mode": args.replacement_mode,
                "skip_candidate_export": args.skip_candidate_export,
                "substitute_reference_path": str(substitute_reference_path),
                "review_scope": (
                    "explicit_substitute_reference_plus_date_blocked_plus_ambiguity"
                    if args.replacement_mode == REPLACEMENT_MODE_TECHNICAL_REFERENCE
                    else "same_custom_code_clean_plus_ambiguity_plus_no_lookup_plus_date_blocked"
                ),
                "commit_scope": (
                    "explicit_substitute_reference_plus_ambiguity_plus_no_lookup"
                    if args.replacement_mode == REPLACEMENT_MODE_TECHNICAL_REFERENCE
                    else "same_custom_code_clean_plus_ambiguity_plus_no_lookup"
                ),
                "future_runtime_selectable": [
                    "bom_selection_strategy",
                    "sequence_strategy",
                    "replacement_pass_scope",
                    "replacement_objective",
                ],
                "hard_coded_defaults": {
                    "starting_points": starting_point_ids,
                    "pass_1_stock_scope": "same_custom_code_clean_plus_ambiguity_plus_no_lookup",
                    "objective": "reach_stock_sufficient_and_rvc_35_then_stop_per_product",
                    "fallback": "maximize_pass_feasibility_then_reduce_non_origin_cost",
                    "bom_output": "per_product_bom_snapshot",
                },
            },
            indent=2,
            ensure_ascii=False,
            default=str,
        )
        + "\n",
        encoding="utf-8",
    )
    write_summary_doc(args.output_doc, args.shipment, policy, seed_results, starting_point_ids)
    return {
        "seed_summaries": seed_results,
        "candidate_csvs": candidate_csvs,
        "review_workbook": workbook_path,
        "workbook_exports": workbook_exports,
        "doc_path": args.output_doc,
    }


def main() -> None:
    run(parse_args())


if __name__ == "__main__":
    main()
